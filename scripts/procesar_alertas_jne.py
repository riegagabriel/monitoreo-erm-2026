"""Alertas del JNE para el dashboard: filtra «RIESGO DE VIOLENCIA ELECTORAL = Sí» y asigna ubigeo INEI.

Fuente (solo lectura): BBDD_JNE/Riesgos_electorales_ERM2026_53casos.xlsx (Fuente: DEE - DNEECT).
Salida: data/alertas_jne_Q<corte>.json, sin datos personales: lugar, JEE y tipo de riesgo.

El cruce con el catalogo INEI usa la tripleta departamento-provincia-distrito. Si la provincia del JNE no
coincide, se acepta departamento + distrito solo cuando es unico en el departamento, y se deja constancia
en `resolucion_territorial`. Un caso que no se pueda ubicar aborta: nunca se descarta en silencio.
"""
from __future__ import annotations

import collections
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "BBDD_JNE" / "Riesgos_electorales_ERM2026_53casos.xlsx"
LOOKUP = RAIZ.parent / "DENUNCIAS_REA" / "data" / "geodata" / "distritos_lookup.parquet"
SALIDA = RAIZ / "data" / "alertas_jne_Q20261001.json"
CORTE = "01/10/2026"  # fecha en que se recibio la base; la base no trae fecha propia


def norm(x) -> str:
    s = unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", s)).strip()


def leer_casos(ruta: Path = FUENTE) -> list[dict]:
    ws = openpyxl.load_workbook(ruta, data_only=True).active
    cab = [str(c).strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    esperado = ["Nº", "DEPARTAMENTO", "JEE", "PROVINCIA", "DISTRITO", "TIPO DE RIESGO ELECTORAL", "RIESGO DE VIOLENCIA ELECTORAL"]
    if cab[:7] != esperado:
        raise SystemExit(f"Cabeceras inesperadas en la base del JNE: {cab[:7]}")
    casos = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None and re.fullmatch(r"\d+(\.0)?", str(r[0]).strip()):
            casos.append({"n": int(float(r[0])), "departamento": str(r[1]).strip(), "jee": str(r[2]).strip(),
                          "provincia": str(r[3]).strip(), "distrito": str(r[4]).strip(),
                          "tipo": str(r[5]).strip(), "violencia": norm(r[6])})
    return casos


def construir(casos: list[dict], catalogo: pd.DataFrame) -> list[dict]:
    triple = collections.defaultdict(list)
    dep_dist = collections.defaultdict(list)
    for r in catalogo.itertuples():
        d, p, t = norm(r.DEPARTAMEN), norm(r.PROVINCIA), norm(r.DISTRITO)
        triple[(d, p, t)].append(r)
        dep_dist[(d, t)].append(r)
    alertas = []
    for c in casos:
        if c["violencia"] != "SI":
            continue
        hit = triple.get((norm(c["departamento"]), norm(c["provincia"]), norm(c["distrito"])), [])
        res = "exacta"
        if len(hit) != 1:
            hit = dep_dist.get((norm(c["departamento"]), norm(c["distrito"])), [])
            res = "por departamento y distrito (la provincia del JNE no coincide con el INEI)"
        if len(hit) != 1:
            raise SystemExit(f"No se puede ubicar el caso {c['n']}: {c['departamento']} / {c['provincia']} / {c['distrito']} ({len(hit)} coincidencias)")
        h = hit[0]
        alertas.append({
            "id": f"JNE-{c['n']:02d}", "n_jne": c["n"], "ubigeo_inei": str(h.UBIGEO_INEI),
            "departamento": str(h.DEPARTAMEN), "provincia": str(h.PROVINCIA), "distrito": str(h.DISTRITO),
            "jee": c["jee"], "tipo_riesgo": c["tipo"], "riesgo_violencia": "Sí",
            "provincia_jne": c["provincia"], "resolucion_territorial": res, "lon": round(float(h.lon), 5), "lat": round(float(h.lat), 5),
        })
    return alertas


def main() -> None:
    casos = leer_casos()
    catalogo = pd.read_parquet(LOOKUP, columns=["UBIGEO_INEI", "DEPARTAMEN", "PROVINCIA", "DISTRITO", "lon", "lat"])
    alertas = construir(casos, catalogo)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps({
        "fuente": "JNE, DEE - DNEECT. Riesgos electorales ERM 2026", "corte": CORTE,
        "total_base": len(casos), "con_riesgo_violencia": len(alertas), "alertas": alertas,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    ajustadas = [a["id"] for a in alertas if a["resolucion_territorial"] != "exacta"]
    print(f"{len(casos)} casos en la base, {len(alertas)} con riesgo de violencia -> {SALIDA}")
    print(f"ubicadas por departamento y distrito: {ajustadas}")
    print("por departamento:", dict(collections.Counter(a["departamento"] for a in alertas).most_common()))


if __name__ == "__main__":
    sys.exit(main())
