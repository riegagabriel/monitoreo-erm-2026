"""Exporta a xlsx todas las bases que alimentan el dashboard, para que una persona no tecnica pueda auditar como se construye.

Uso: python scripts/exportar_auditoria.py            (usa los archivos de data/ y data/entrada/ tal como estan; corra antes corte.py y matriz_a_dashboard.py --aplicar)
Salida: seguimiento/AUDITORIA_DASHBOARD_Q<aaaammdd>_<hhmm>/  (carpeta ignorada por Git: lleva nombres de orientadores, uso interno)

Archivos: 00 guia y KPIs, 01 incidencias (central), 02 marcadores del mapa, 03 Forms y orientadores, 04 consultas, 05 distritos y alertas previas.
Los marcadores y el orden de la tabla se recalculan aqui con las mismas reglas del dashboard (ultima version de bocetos/dashboard_erm_vN.html).
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import corte as c  # noqa: E402
from actualizar_locales import BASE, HOJA, nz  # noqa: E402

R = c.RAIZ
VIOL = re.compile(r"\b(violen|amenaz|agresi|agred|golpe|pelea|rina\b|enfrentamiento|disturbio|atac|empujon|incendi|quem[oa]|destrozo|balacera|disparo|herid|lesion|intimid|hostig)")
CANALES = {"ori": "Orientadores", "rea": "Denuncias REA de hoy", "ent": "Entidades del sistema", "pre": "Prensa y redes sociales", "tel": "Llamadas o mesa de ayuda"}
UMB = [100, 500, 3000]
TRAMOS = ["1–99", "100–499", "500–2 999", "≥ 3 000"]
TIPOS_CONSULTA = ["DNI vencido", "Restitución de domicilio", "Cara de niño", "Ciudadanos fallecidos", "Otras consultas"]
DIA = dt.datetime(2026, 10, 4)


# ---------------------------------------------------------------- utilidades
def norm(s) -> str:
    return unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()


def grave(x: dict) -> bool:
    return x["g"] == 2 or bool(VIOL.search(norm(x["t"])))


def hh(h: float) -> str:
    return f"{int(h):02d}:{round((h % 1) * 60):02d}"


def fn(n) -> str:
    return f"{int(n):,}".replace(",", " ")


def enmascara(t) -> str:
    """Los textos originales son de uso interno, pero igual se tapan DNI (8 digitos) y telefonos (9 digitos)."""
    return re.sub(r"\b\d{9}\b", "[teléfono omitido]", re.sub(r"\b\d{8}\b", "[DNI omitido]", str(t or "")))


def titulo(s) -> str:
    return " ".join(w.capitalize() for w in str(s or "").split())


def cargar(ruta: Path):
    return json.loads(ruta.read_text(encoding="utf-8"))


def csv_por(ruta: Path, llave: str) -> dict[str, dict]:
    if not ruta.exists():
        return {}
    with ruta.open(encoding="utf-8", newline="") as f:
        return {x[llave].strip(): x for x in csv.DictReader(f)}


def guarda(carpeta: Path, nombre: str, hojas: list[tuple]) -> Path:
    """hojas: (titulo, columnas, filas, anchos opcionales {col: ancho}). La primera columna con texto largo se ajusta sola."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tit, cols, filas, *ancho in hojas:
        ws = wb.create_sheet(tit[:31])
        ancho = ancho[0] if ancho else {}
        ws.append(cols)
        for f in filas:
            ws.append([("" if v is None else v) for v in f])
        for j, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=j)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            largo = max([len(str(col))] + [len(str(ws.cell(row=i, column=j).value or "")) for i in range(2, min(ws.max_row, 60) + 1)])
            ws.column_dimensions[get_column_letter(j)].width = ancho.get(col, min(max(11, largo + 2), 70))
        for fila in ws.iter_rows(min_row=2):
            for cell in fila:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        if filas:
            ws.auto_filter.ref = ws.dimensions
    ruta = carpeta / nombre
    wb.save(ruta)
    return ruta


def leeme(lineas: list[tuple[str, str]]) -> tuple:
    return ("LEEME", ["Tema", "Explicación"], lineas, {"Tema": 34, "Explicación": 120})


# ---------------------------------------------------------------- carga de fuentes
def ultimo_boceto() -> Path:
    return sorted(R.glob("bocetos/dashboard_erm_v*.html"), key=lambda p: int(re.search(r"_v(\d+)", p.stem).group(1)))[-1]


HTML = ultimo_boceto()
TXT = HTML.read_text(encoding="utf-8")
ARCH = {k: re.search(rf"{k}='\.\./data/([^']+)'", TXT).group(1) for k in ("ORIE", "JNE_", "BAN_", "COORD", "JOR")}
J = cargar(R / "data" / ARCH["JOR"])
BANJ = cargar(R / "data" / ARCH["BAN_"])
JNEJ = cargar(R / "data" / ARCH["JNE_"])
ORIE = cargar(R / "data" / ARCH["ORIE"])
COORDJ = cargar(R / "data" / ARCH["COORD"])
PRE = R / "dashboard" / "public" / "data" / "precarga"
RIESGO = cargar(PRE / "riesgo_previo.json")
GEO = cargar(PRE / "distritos_pais.geojson")
IDX = {f["properties"]["u"]: f["properties"] for f in GEO["features"]}
RISK = {d["ubigeo_inei"]: d for d in RIESGO["distritos"]}
LOCALES = {r["id"]: r for r in ORIE["registros"]}
COORD = {r["id"]: r for r in COORDJ["registros"]}


def rest_de(u: str) -> dict:
    return (RISK.get(u) or {}).get("restituidos") or {"total": 0, "reniec": 0, "jne": 0}


def ubic(u: str) -> tuple:
    p = IDX.get(u) or {}
    return p.get("dep", ""), p.get("prov", ""), p.get("dist", "")


# ---------------------------------------------------------------- Forms
FORMS = R / "data" / "entrada" / "hoja_respuestas.xlsx"
filas_forms = list(openpyxl.load_workbook(FORMS, data_only=True)[c.PESTANA].iter_rows(values_only=True))
REG = c.normalizar(list(filas_forms[0]), filas_forms[1:])
BASE_LIDS = c.cargar_base(ORIE)
POR_LOCAL = {(nz(r["distrito"]), nz(r["local"])): r["id"] for r in ORIE["registros"]}
INICIO = c.INICIO_JORNADA


def leer_base_vf() -> dict[str, dict]:
    ws = openpyxl.load_workbook(BASE, data_only=True)[HOJA]
    cab = [nz(x) if x else "" for x in next(ws.iter_rows(values_only=True))]
    ix = {}
    for i, h in enumerate(cab):
        ix.setdefault(h, i)
    out: dict[str, dict] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[ix["NOMBRES Y APELLIDOS"]]:
            continue
        if r[ix["N"]] not in (None, "") and not str(r[ix["N"]]).strip().isdigit():
            continue
        nombre = " ".join(str(r[ix["NOMBRES Y APELLIDOS"]]).split())
        x = out.setdefault(nz(nombre), {"nombre": nombre, "monitor": "", "oficina": str(r[ix["OFICINA REGIONAL"]] or "").strip(), "locales": []})
        x["monitor"] = x["monitor"] or " ".join(str(r[ix["MONITOR OPERATIVO"]] or "").split())
        lid = POR_LOCAL.get((nz(r[ix["DISTRITO"]]), nz(r[ix["NOMBRE DEL LOCAL"]])))
        if lid and lid not in x["locales"]:
            x["locales"].append(lid)
    return out


BASEV = leer_base_vf()
NOMBRE = {k: v["nombre"] for k, v in BASEV.items()}


def lids_de(r: dict) -> list[str]:
    return c._lids(r, BASE_LIDS, POR_LOCAL)


def id_forms(r: dict):
    lid = POR_LOCAL.get((r["dist"], r["local"]))
    return f"ORI-{r['ts']:%Y%m%d%H%M%S}-{lid}" if lid else None


INC_FORMS = csv_por(R / "data" / "entrada" / "incidencias_forms.csv", "id_envio")
FORMS_INC = {id_forms(r): r for r in REG if r["ts"] >= INICIO and r["inc"] == "Sí" and id_forms(r)}

# ---------------------------------------------------------------- matriz
MATRIZ = R / "data" / "entrada" / "matriz_casos.xlsx"
VALID = csv_por(R / "data" / "entrada" / "validaciones_matriz.csv", "id")
ESTADO = csv_por(R / "data" / "entrada" / "estado_matriz.csv", "id")
PREF = {"id": "ID", "dep": "DEPARTAMENTO", "prov": "PROVINCIA", "dist": "DISTRITO", "ubigeo": "UBIGEO RENIEC", "local": "LOCAL DE VOTACION", "fecha": "FECHA REPORTE", "hora": "HORA",
        "remitente": "REMITENTE", "fuente": "FUENTE ORIGINAL", "tipo_doc": "TIPO DE DOCUMENTO", "categoria": "CATEGORIA", "desc": "DESCRIPCION RESUMIDA", "acciones": "ACCIONES",
        "gravedad": "GRAVEDAD", "bandera": "BANDERA ROJA", "adjunto": "ADJUNTO", "obs": "OBSERVACIONES", "responsable": "RESPONSABLE", "registrado": "SE HA REGISTRADO"}


def leer_matriz() -> list[dict]:
    ws = openpyxl.load_workbook(MATRIZ, data_only=True)["Casos"]
    filas = list(ws.iter_rows(values_only=True))
    cab = [nz(x) for x in filas[0]]
    ix = {}
    for k, p in PREF.items():
        ix[k] = next((i for i, h in enumerate(cab) if (h == p if k == "id" else h.startswith(p))), None)
    out, vistos = [], Counter()
    for n, r in enumerate(filas[1:], start=2):
        if r[ix["id"]] in (None, ""):
            continue
        d = {k: (r[i] if i is not None else None) for k, i in ix.items()}
        d["fila"] = n
        cid = str(d["id"]).strip()
        vistos[cid] += 1
        d["id_hoja"] = cid
        d["id"] = cid if vistos[cid] == 1 else f"{cid}-{vistos[cid]}"
        out.append(d)
    return out


CASOS = leer_matriz()
CASO = {x["id"]: x for x in CASOS}


def caso_de(did: str):
    k = did[4:]
    return CASO.get(k) or CASO.get(re.sub(r"-\d+$", "", k))


def ids_dash(cid: str, universo) -> list[str]:
    """Ids del dashboard (MAT-C07, MAT-C07-1...) que vienen del caso cid; «MAT-C25-2» es del caso C25-2, no del C25."""
    out = []
    for i in universo:
        if not str(i).startswith("MAT-"):
            continue
        k = i[4:]
        if k == cid or (re.fullmatch(rf"{re.escape(cid)}-\d+", k) and k not in CASO):
            out.append(i)
    return out


def estado_sin_publicar(cs: dict) -> tuple[bool, str]:
    """(cubierto por una bandera REA del distrito, explicacion) para un caso que no tiene registro propio en el dashboard."""
    e, v = ESTADO.get(cs["id"], {}), VALID.get(cs["id"], {})
    det = e.get("detalle", "")
    if "ya tiene bandera" in det:
        return True, det
    if "(MAT-" in det:  # estado_matriz.csv confunde ids repetidos de la hoja: se explica con la revision
        det = v.get("motivo") or "Aún sin validar para el dashboard (falta la validación del equipo)"
    return False, det


def fecha_txt(v) -> str:
    return v.strftime("%d/%m/%Y") if isinstance(v, (dt.datetime, dt.date)) else str(v or "")


def hora_txt(v) -> str:
    return v.strftime("%H:%M") if isinstance(v, (dt.datetime, dt.time)) else str(v or "")


# ---------------------------------------------------------------- lo que dibuja el dashboard (mismas reglas que el HTML)
INC = J["incidencias"]
BAN = [dict(x, u=x["ubigeo_inei"]) for x in BANJ["registros"] if x["conflictividad"]]
EN_BAN = {int(b["id"][4:]) for b in BAN if b["id"].startswith("REA-")}
DEN = list(BAN)
for d in RIESGO["distritos"]:
    for a in d["alertas_rea"]:
        if a["item"] not in EN_BAN:
            DEN.append({"id": f"REA-{a['item']}", "u": d["ubigeo_inei"], "resumen_publicable": a["motivo"], "precarga": True, "conflictividad": None})
JNE = [dict(a, u=a["ubigeo_inei"]) for a in JNEJ["alertas"]]


def marcadores() -> tuple[list[dict], list[dict]]:
    """(marcadores, registros por marcador). Orientacion y alertas del mapa principal; locales del mapa de la pestana Locales."""
    mk, rg = [], []
    vistos: Counter = Counter()
    for l in ORIE["registros"]:
        vistos[l["ubigeo_inei"]] += 1
        k = vistos[l["ubigeo_inei"]] - 1
        dx = 0 if not k else (12 if k % 2 else -12)
        p = IDX[l["ubigeo_inei"]]
        mid = f"MK-ORI-{l['id']}"
        mk.append({"id": mid, "capa": "Orientación RENIEC (mapa principal)", "simbolo": "Pin azul con el n.º de orientadores", "pestana": "Mapa y alertas", "u": l["ubigeo_inei"], "lon": p["lon"], "lat": p["lat"],
                   "origen_pos": "Centro del distrito (geojson de distritos)" + (f"; desplazado {dx} px porque hay otro local en el distrito" if dx else ""), "n": 1, "ids": l["id"],
                   "rotulo": f"Orientación electoral · {l['local']} · {l['distrito']}", "interruptor": "Orientación", "datos": f"data/{ARCH['ORIE']} (registro {l['id']})", "obs": f"{l['n_orientadores']} orientadores asignados"})
        rg.append({"marcador": mid, "tipo": "Local de orientación", "id": l["id"], "texto": f"{l['local']} · {l['n_orientadores']} orientadores"})
    for a in JNE:
        p = IDX[a["u"]]
        mid = f"MK-{a['id']}"
        mk.append({"id": mid, "capa": "Alerta del JNE (mapa principal)", "simbolo": "Logo del JNE", "pestana": "Mapa y alertas", "u": a["u"], "lon": p["lon"], "lat": p["lat"], "origen_pos": "Centro del distrito (geojson de distritos)",
                   "n": 1, "ids": a["id"], "rotulo": "Alerta del JNE", "interruptor": "Alertas JNE", "datos": f"data/{ARCH['JNE_']} (alerta {a['id']})", "obs": a["tipo_riesgo"]})
        rg.append({"marcador": mid, "tipo": "Alerta del JNE", "id": a["id"], "texto": a["tipo_riesgo"]})
    por_u: dict[str, list] = defaultdict(list)
    for b in DEN:
        por_u[b["u"]].append(b)
    for u, bs in por_u.items():
        p = IDX[u]
        mid = f"MK-BR-{u}"
        mk.append({"id": mid, "capa": "Alerta RENIEC (mapa principal)", "simbolo": "Bandera roja", "pestana": "Mapa y alertas", "u": u, "lon": p["lon"], "lat": p["lat"], "origen_pos": "Centro del distrito; la bandera sale a la izquierda",
                   "n": len(bs), "ids": ", ".join(b["id"] for b in bs), "rotulo": "Alerta RENIEC" + (f" ({len(bs)})" if len(bs) > 1 else ""), "interruptor": "Alertas RENIEC",
                   "datos": f"data/{ARCH['BAN_']} (conflictividad = Sí) + alertas REA de precarga/riesgo_previo.json", "obs": "Una bandera por distrito; el número entre paréntesis es la cantidad de alertas"})
        for b in bs:
            rg.append({"marcador": mid, "tipo": "Alerta RENIEC", "id": b["id"], "texto": b["resumen_publicable"]})
    graves: dict[str, list] = defaultdict(list)
    for x in INC:
        if grave(x):
            graves[x["u"]].append(x)
    for u, xs in graves.items():
        p = IDX[u]
        mid = f"MK-BV-{u}"
        mk.append({"id": mid, "capa": "Incidencia grave del domingo (mapa principal)", "simbolo": "Bandera verde", "pestana": "Mapa y alertas", "u": u, "lon": p["lon"], "lat": p["lat"], "origen_pos": "Centro del distrito; la bandera sale a la derecha",
                   "n": len(xs), "ids": ", ".join(x["id"] for x in xs), "rotulo": "Incidencia grave del domingo" + (f" ({len(xs)})" if len(xs) > 1 else ""), "interruptor": "Incidencias",
                   "datos": "data/jornada.json (incidencias con g = 2 o texto de violencia)", "obs": "Solo las graves se dibujan; las demás cuentan en cifras, gráficos y tablas"})
        for x in xs:
            rg.append({"marcador": mid, "tipo": "Incidencia grave", "id": x["id"], "texto": x["t"]})
    for l in ORIE["registros"]:
        cc = COORD.get(l["id"]) or {}
        con = cc.get("lat") is not None
        p = IDX[l["ubigeo_inei"]]
        mid = f"MK-LOC-{l['id']}"
        mk.append({"id": mid, "capa": "Local de votación (pestaña Locales)", "simbolo": "Pin azul con el n.º de orientadores", "pestana": "Locales de orientación", "u": l["ubigeo_inei"],
                   "lon": cc["lon"] if con else p["lon"], "lat": cc["lat"] if con else p["lat"],
                   "origen_pos": f"Coordenada del colegio (padrón MINEDU, confianza «{cc.get('confianza')}»; código de local {cc.get('codlocal')})" if con else "Centro del distrito (el colegio no tiene coordenada)",
                   "n": 1, "ids": l["id"], "rotulo": f"{l['local']} · {l['distrito']}", "interruptor": "—", "datos": f"data/{ARCH['COORD']} (registro {l['id']})", "obs": "Pin punteado si no hay coordenada" if not con else ""})
        rg.append({"marcador": mid, "tipo": "Local de orientación", "id": l["id"], "texto": f"{l['local']}"})
    return mk, rg


MK, MKREG = marcadores()
MK_DE = {}
for r in MKREG:
    if r["marcador"].startswith(("MK-BR", "MK-BV", "MK-JNE")):
        MK_DE[r["id"]] = r["marcador"]


# ---------------------------------------------------------------- orden de la tabla del dashboard
def tabla_dashboard() -> list[dict]:
    todas = []
    for x in INC:
        todas.append({"t": "dom", "u": x["u"], "h": x["h"], "prio": 0 if grave(x) else 2, "x": x})
    for b in DEN:
        todas.append({"t": "ren", "u": b["u"], "h": 0, "prio": 1 if b.get("conflictividad") else 3, "x": b})
    for a in JNE:
        todas.append({"t": "jne", "u": a["u"], "h": 0, "prio": 4, "x": a})
    for i, r in enumerate(todas):
        r["i"] = i
    todas.sort(key=lambda r: (r["prio"], -(r["h"] or 0), r["i"]))
    for n, r in enumerate(todas, start=1):
        r["orden"] = n
    return todas


TABLA = tabla_dashboard()
PRIO = {0: "Grave", 1: "Riesgo de conflicto", 2: "Atención", 3: "Alerta previa sin bandera", 4: "Alerta JNE previa"}


def razon_grave(x: dict) -> str:
    r = []
    if x["g"] == 2:
        r.append("el equipo la etiquetó como grave (tipo A al validar el Forms)" if x["id"].startswith("ORI-") else "la hoja de la matriz la marca con gravedad «Alta»")
    m = VIOL.search(norm(x["t"]))
    if m:
        r.append(f"el texto menciona «{m.group(0)}…» (palabra de violencia)")
    return " y ".join(r) if r else "No es grave: se muestra en tablas y cifras, no en el mapa"


# ---------------------------------------------------------------- 01 incidencias central
def fila_comun(r: dict) -> dict:
    u = r["u"]
    dep, prov, dist = ubic(u)
    rs = rest_de(u)
    return {"orden": r["orden"], "tipo": {"dom": "Incidencia (domingo 4)", "ren": "Alerta RENIEC", "jne": "Alerta del JNE"}[r["t"]], "prio": PRIO[r["prio"]], "dep": dep, "prov": prov, "dist": dist, "u": u,
            "rest": rs["total"], "rest_reniec": rs["reniec"], "rest_jne": rs["jne"]}


def central() -> list[list]:
    out = []
    for r in TABLA:
        x, f = r["x"], fila_comun(r)
        local = ""
        obs, fuente, donde, idf, orig, quien, revision, cambios, ruta = [], "", "", "", "", "", "", "", ""
        fecha = hora = canal = ""
        if r["t"] == "dom":
            fecha, hora = "04/10/2026", hh(x["h"])
            canal = CANALES.get(x["c"], x["c"])
            if x["l"]:
                local = f"{x['l']} · {LOCALES[x['l']]['local']}"
            if x["id"].startswith("ORI-"):
                fr = FORMS_INC.get(x["id"]) or {}
                cv = INC_FORMS.get(x["id"], {})
                fuente = "Formulario de orientadores (Google Forms), incidencia que reportó el orientador"
                donde = f"Google Forms · hoja RESP_ORIENTADORES · fila {fr.get('fila', '?')} (envío del {fr['ts']:%d/%m/%Y %H:%M:%S})" if fr else "Google Forms (no se encontró la fila)"
                idf = x["id"]
                orig = enmascara(fr.get("det"))
                quien = f"Orientador: {titulo(NOMBRE.get(fr.get('k'), fr.get('k')))}" if fr else ""
                revision = f"Revisión humana en incidencias_forms.csv: validado={cv.get('validado', '')}; tipo={cv.get('tipo', '')} ({'A = grave' if cv.get('tipo') == 'A' else 'B = de atención'}); nota: {cv.get('nota', '')}"
                cambios = "Sin cambios" if enmascara(fr.get("det")).strip() == x["t"].strip() else "El equipo redactó un resumen sin nombres ni datos personales (≤ 280 caracteres)"
                ruta = "Orientador responde el Forms → el equipo revisa y redacta el resumen (incidencias_forms.csv) → corte.py lo pasa a jornada.json → el dashboard lo lee"
            else:
                cs = caso_de(x["id"])
                v = VALID.get(cs["id"], {}) if cs else {}
                fuente = "Matriz de sistematización de denuncias (hoja Casos del equipo), canal de bandera verde"
                donde = f"Google Sheet «Casos» · fila {cs['fila']} · ID {cs['id_hoja']}" if cs else "Matriz (caso no encontrado)"
                idf = cs["id_hoja"] if cs else x["id"]
                orig = enmascara(cs["desc"]) if cs else ""
                quien = f"Remitente: {cs['remitente']}; responsable de verificar: {cs['responsable']}" if cs else ""
                revision = f"Revisión humana en validaciones_matriz.csv: validado={v.get('validado', '')}; canal={v.get('canal', '')}; motivo: {v.get('motivo', '')}"
                cambios = "Texto reemplazado por el de la revisión (el de la hoja no se usó tal cual)" if v.get("texto") else "Texto tomado de la hoja"
                if cs and v.get("fecha"):
                    cambios += "; fecha corregida en la revisión"
                if cs:
                    obs.append(f"Gravedad en la hoja: {cs['gravedad']}; bandera roja RENIEC: {cs['bandera']}; fuente original: {cs['fuente']} · {cs['tipo_doc']}; adjuntos: {cs['adjunto'] or '—'}")
                ruta = "El equipo llena la matriz → matriz_a_dashboard.py la evalúa (fecha, ubicación, texto) y la pasa a incidencias_matriz.json → corte.py la suma a jornada.json → el dashboard la lee"
            grave_ = grave(x)
            f.update(fecha=fecha, hora=hora, canal=canal)
            fila = [r["orden"], x["id"], f["tipo"], f["prio"], "Sí" if grave_ else "No (solo las graves se dibujan)", f"MK-BV-{x['u']}" if grave_ else "—",
                    "Bandera verde" if grave_ else "—", f["dep"], f["prov"], f["dist"], x["u"], local, fecha, hora, canal, x["t"], razon_grave(x), fuente, donde, idf, orig, quien, revision, cambios,
                    f"data/jornada.json → incidencias[ id = {x['id']} ]", ruta,
                    "Tabla «Detalle de las alertas e incidencias»" + (" (filtro Graves)" if grave_ else "") + (" · Mapa: bandera verde" if grave_ else "") + " · tarjeta «incidencias del domingo» · gráficos por hora, por canal y previas/domingo"
                    + (" · pestaña Locales (columna incidencias)" if x["l"] else ""), f["rest"], f["rest_reniec"], f["rest_jne"], " | ".join(obs)]
        elif r["t"] == "ren":
            fecha = (x.get("fecha_recepcion") or "—") + (" (asumida)" if x.get("fecha_asumida") else "")
            canal = x.get("canal") or "Precargado REA"
            idf = x["id"]
            if x["id"].startswith("MAT-"):
                cs = caso_de(x["id"])
                v = VALID.get(cs["id"], {}) if cs else {}
                fuente = "Matriz de sistematización de denuncias (hoja Casos del equipo), canal de bandera roja"
                donde = f"Google Sheet «Casos» · fila {cs['fila']} · ID {cs['id_hoja']}" if cs else "Matriz"
                orig = enmascara(cs["desc"]) if cs else ""
                quien = f"Remitente: {cs['remitente']}; responsable de verificar: {cs['responsable']}" if cs else ""
                revision = f"Revisión humana en validaciones_matriz.csv: validado={v.get('validado', '')}; canal={v.get('canal', '')}; motivo de la bandera: {x.get('motivo_conflictividad', '')}"
                cambios = "Texto reemplazado por el de la revisión" if v.get("texto") else "Texto tomado de la hoja"
                ruta = "El equipo llena la matriz (bandera roja = Sí) → matriz_a_dashboard.py --aplicar la agrega a banderas_rea → el dashboard la lee"
                if cs:
                    obs.append(f"Gravedad en la hoja: {cs['gravedad']}; fuente original: {cs['fuente']} · {cs['tipo_doc']}; adjuntos: {cs['adjunto'] or '—'}")
                idf = f"{x['id']} (hoja: {cs['id_hoja']})" if cs else x["id"]
            elif x.get("precarga"):
                fuente = "REA precargado (alertas previas con denuncia, casos.json corte 30/09/2026)"
                donde = "dashboard/public/data/precarga/riesgo_previo.json → distritos[].alertas_rea"
                orig = f"Motivo de la alerta REA ítem {x['id'][4:]} (texto del precargado)"
                revision = "Precargado antes de la jornada; sin revisión nueva"
                cambios = "Sin cambios"
                ruta = "Base REA → precarga riesgo_previo.json → el dashboard la suma a las alertas RENIEC si no está en banderas_rea"
            else:
                fuente = "Registro de Alertas Electorales (REA): denuncia informada directamente al equipo"
                donde = f"Base consolidada REA, ítem {x['item']} → data/{ARCH['BAN_']} (id {x['id']})"
                orig = "Observación del ítem en la base REA (DENUNCIAS_REA); no se copia aquí porque puede traer datos personales"
                quien = f"Entidad emisora: {x.get('entidad_emisora') or '—'}"
                revision = f"Estado: {x.get('estado')}; categoría REA {x.get('categoria_rea')}; riesgo de conflicto: Sí; motivo: {x.get('motivo_conflictividad', '')}"
                cambios = "El equipo redactó un resumen publicable sin nombres de ciudadanos"
                ruta = "Denuncia o correo → base REA → script 05 de DENUNCIAS_REA → banderas_rea_Q… .json → el dashboard lo lee"
                if x.get("n_ciudadanos"):
                    obs.append(f"Ciudadanos listados en la denuncia: {x['n_ciudadanos']}")
            f.update(fecha=fecha, hora="", canal=canal)
            fila = [r["orden"], x["id"], f["tipo"], f["prio"], "Sí", f"MK-BR-{x['u']}", "Bandera roja", f["dep"], f["prov"], f["dist"], x["u"], "", fecha, "", canal, x["resumen_publicable"], "No aplica (alerta previa)", fuente,
                    donde, idf, orig, quien, revision, cambios, f"data/{ARCH['BAN_']} → registros[ id = {x['id']} ]" if not x.get("precarga") else "precarga/riesgo_previo.json",
                    ruta, "Tabla «Detalle…» (filtros Todas y Alertas RENIEC) · Mapa: bandera roja del distrito · tarjeta «alertas RENIEC» · gráfico «Previas y del domingo» · pestaña Locales (columna alertas previas)",
                    f["rest"], f["rest_reniec"], f["rest_jne"], " | ".join(obs)]
        else:
            fila = [r["orden"], x["id"], f["tipo"], f["prio"], "Sí", f"MK-{x['id']}", "Logo del JNE", f["dep"], f["prov"], f["dist"], x["u"], "", "—", "", f"JNE · JEE {x['jee']}",
                    "Riesgo de violencia electoral. " + x["tipo_riesgo"], "No aplica (alerta previa)", "Base de riesgos del JNE (JEE) enviada antes de la jornada", f"data/{ARCH['JNE_']} (alerta {x['id']}, JEE {x['jee']})", x["id"],
                    x["tipo_riesgo"], "", f"Riesgo de violencia: {x['riesgo_violencia']}", "Sin cambios", f"data/{ARCH['JNE_']} → alertas[ id = {x['id']} ]",
                    "JEE → base JNE → equipo de datos empareja el distrito → alertas_jne_Q… .json → el dashboard lo lee",
                    "Tabla «Detalle…» (filtros Todas y Alertas del JNE) · Mapa: logo JNE · tarjeta «alertas del JNE» · gráfico «Previas y del domingo»", f["rest"], f["rest_reniec"], f["rest_jne"],
                    f"Resolución territorial: {x['resolucion_territorial']}"]
        out.append(fila)
    return out


COLS_CENTRAL = ["N.° EN LA TABLA", "ID DEL REGISTRO", "TIPO EN LA TABLA", "PRIORIDAD (orden de la tabla)", "¿SE DIBUJA EN EL MAPA?", "ID DEL MARCADOR EN EL MAPA", "SÍMBOLO EN EL MAPA", "DEPARTAMENTO", "PROVINCIA",
                "DISTRITO", "UBIGEO INEI", "LOCAL DE VOTACIÓN", "FECHA", "HORA", "FUENTE QUE MUESTRA LA TABLA", "TEXTO PUBLICADO EN EL DASHBOARD", "¿POR QUÉ ES GRAVE?", "FUENTE DE ORIGEN", "DÓNDE ESTÁ EL DATO ORIGINAL",
                "ID EN LA FUENTE", "TEXTO ORIGINAL EN LA FUENTE", "QUIÉN LO REPORTÓ (uso interno)", "REVISIÓN HUMANA", "CAMBIOS ENTRE EL ORIGINAL Y LO PUBLICADO", "ARCHIVO Y CAMPO QUE LEE EL DASHBOARD", "CÓMO LLEGÓ AL DASHBOARD",
                "DÓNDE SE VE (tablas, mapa, tarjetas)", "RESTITUIDOS EN EL DISTRITO (total)", "RESTITUIDOS RENIEC", "RESTITUIDOS JNE", "OBSERVACIONES Y AVISOS"]


def no_publicadas() -> list[list]:
    out = []
    publicados = {x["id"] for x in INC} | {b["id"] for b in BANJ["registros"]}
    for cs in CASOS:
        if ids_dash(cs["id"], publicados):
            continue
        e, v = ESTADO.get(cs["id"], {}), VALID.get(cs["id"], {})
        dep, prov, dist = cs["dep"], cs["prov"], cs["dist"]
        cub, det = estado_sin_publicar(cs)
        estado = "Cubierto por una bandera REA del mismo distrito (no se publica aparte)" if cub else "No publicado"
        out.append([f"Matriz · {cs['id']}", cs["id_hoja"], f"fila {cs['fila']}", estado, det, f"validado={v.get('validado', '—')}; motivo: {v.get('motivo', '')}", dep, prov, dist, cs["ubigeo"], fecha_txt(cs["fecha"]), hora_txt(cs["hora"]),
                    cs["gravedad"], cs["bandera"], enmascara(cs["desc"]), cs["remitente"], cs["responsable"], f"{cs['fuente']} · {cs['tipo_doc']} · {cs['adjunto'] or '—'}", cs["registrado"]])
    for ide, r in FORMS_INC.items():
        if ide in {x["id"] for x in INC}:
            continue
        cv = INC_FORMS.get(ide)
        out.append(["Forms · incidencia del orientador", ide, f"fila {r['fila']}", "No publicado" if cv else "Pendiente de validar", f"decisión del equipo: {cv['validado'] if cv else 'sin revisar'}", cv.get("nota", "") if cv else "",
                    "", "", r["dist"], "", f"{r['ts']:%d/%m/%Y}", f"{r['ts']:%H:%M}", "", "", enmascara(r["det"]), titulo(NOMBRE.get(r["k"], r["k"])), "", "Google Forms · RESP_ORIENTADORES", ""])
    for b in BANJ["registros"]:
        if not b["conflictividad"]:
            out.append([f"REA · {b['id']}", b["id"], f"ítem {b['item']}", "En el archivo, pero sin riesgo de conflicto: no se dibuja ni aparece en las tablas", "conflictividad = No", b.get("motivo_conflictividad") or "",
                        b["departamento"], b["provincia"], b["distrito"], b["ubigeo_inei"], b["fecha_recepcion"], "", "", "", b["resumen_publicable"], "", "", b.get("canal") or "", ""])
    return out


COLS_NOPUB = ["ORIGEN", "ID", "UBICACIÓN EN LA FUENTE", "ESTADO", "POR QUÉ NO ESTÁ (o cómo está cubierto)", "DECISIÓN / NOTAS", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "UBIGEO", "FECHA", "HORA", "GRAVEDAD EN LA HOJA",
              "BANDERA ROJA EN LA HOJA", "TEXTO ORIGINAL", "REMITENTE / ORIENTADOR", "RESPONSABLE", "FUENTE ORIGINAL", "¿SE HA REGISTRADO EN DASHBOARD? (lo que dice la hoja)"]


def matriz_vs_dashboard() -> list[list]:
    out = []
    ids_inc = {x["id"] for x in INC}
    ids_ban = {b["id"] for b in BANJ["registros"]}
    for cs in CASOS:
        e, v = ESTADO.get(cs["id"], {}), VALID.get(cs["id"], {})
        verde = ids_dash(cs["id"], ids_inc)
        roja = ids_dash(cs["id"], ids_ban)
        cub, det = estado_sin_publicar(cs) if not (verde or roja) else (False, "Publicado como " + ", ".join(verde + roja))
        destino = ("Incidencia del domingo (bandera verde si es grave)" if verde else "Alerta RENIEC (bandera roja)" if roja else
                   "Cubierto por bandera REA del distrito" if cub else "No está en el dashboard")
        out.append([cs["id_hoja"] if cs["id"] == cs["id_hoja"] else f"{cs['id_hoja']} (repetido en la hoja, 2.ª fila)", cs["fila"], cs["dist"], cs["prov"], cs["dep"], fecha_txt(cs["fecha"]), cs["bandera"], destino, ", ".join(verde + roja) or "—",
                    cs["registrado"], e.get("valor_sugerido", ""), e.get("coincide", ""), det, f"validado={v.get('validado', '—')}"])
    return out


# ---------------------------------------------------------------- 03 Forms y orientadores, 04 consultas
def mejores_consultas() -> dict[tuple, dict]:
    mejor: dict[tuple, dict] = {}
    for r in REG:
        if r["ts"] >= INICIO and r["corte"] in (2, 3) and r.get("consultas"):
            k = (r["k"], r["corte"])
            if k not in mejor or sum(r["consultas"].values()) > sum(mejor[k]["consultas"].values()):
                mejor[k] = r
    return mejor


MEJOR = mejores_consultas()


def envios() -> list[list]:
    out = []
    for r in REG:
        prueba = r["ts"] < INICIO
        ide = id_forms(r) if r["inc"] == "Sí" and not prueba else None
        cv = INC_FORMS.get(ide) if ide else None
        publicada = ide in {x["id"] for x in INC} if ide else False
        cons = r.get("consultas") or {}
        usa = bool(cons) and r["corte"] == 3 and not prueba and MEJOR.get((r["k"], r["corte"])) is r
        out.append([r["fila"], f"{r['ts']:%d/%m/%Y %H:%M:%S}", titulo(NOMBRE.get(r["k"], r["k"])), ", ".join(lids_de(r)) if not prueba else "", r["dist"].title(), r["local"].title(), f"Corte {r['corte']}",
                    hora_txt(r["llegada"]), hora_txt(r["termino"]), "PRUEBA (antes del 04/10): no se usa" if prueba else "Real",
                    "Cuenta como llegada" if r["corte"] == 1 and r["llegada"] and not prueba else "Cuenta como cierre" if r["corte"] == 3 and r["termino"] and not prueba else "Corte 2: cuenta en «registros del corte 2»" if r["corte"] == 2 and not prueba else "—",
                    r["inc"] or "", enmascara(r["det"]) if r["inc"] == "Sí" else "", ide or "", ("Publicada" if publicada else (f"No publicada (validado={cv['validado']})" if cv else "Pendiente de validar")) if ide else "",
                    sum(cons.values()) if cons else "", "Sí" if usa else ("No: el criterio usa solo el corte 3" if cons and not prueba and r["corte"] == 2 else "No: reenvío con menor total" if cons and not prueba else "—")])
    return out


def orientadores() -> list[list]:
    out = []
    for k, b in BASEV.items():
        rs = [r for r in REG if r["k"] == k and r["ts"] >= INICIO]
        ll = [r["llegada"] for r in rs if r["corte"] == 1 and r["llegada"]]
        cie = [r["termino"] for r in rs if r["corte"] == 3 and r["termino"]]
        c2 = [r for r in rs if r["corte"] == 2]
        out.append([b["nombre"].title(), b["monitor"].title(), b["oficina"], ", ".join(b["locales"]), " / ".join(LOCALES[l]["local"] for l in b["locales"]), LOCALES[b["locales"][0]]["distrito"], LOCALES[b["locales"][0]]["departamento"],
                    min(ll).strftime("%H:%M") if ll else "Falta", f"{min(r['ts'] for r in c2):%H:%M}" if c2 else "Falta", min(cie).strftime("%H:%M") if cie else "Falta (aún)", len(rs),
                    sum(sum(v["consultas"].values()) for (kk, cc), v in MEJOR.items() if kk == k and cc == 3)])
    return out


def locales_fila() -> list[list]:
    out = []
    av = J["avance"]
    for l in ORIE["registros"]:
        a = av.get(l["id"], {"llegaron": 0, "cierre": 0})
        cc = COORD.get(l["id"]) or {}
        rs = rest_de(l["ubigeo_inei"])
        r2 = {r["k"] for r in REG if r["ts"] >= INICIO and r["corte"] == 2 and l["id"] in lids_de(r)}
        out.append([l["id"], l["departamento"], l["provincia"], l["distrito"], l["ubigeo_inei"], l["local"], l["n_orientadores"], a["llegaron"], f"{a['llegaron']}/{l['n_orientadores']}", len(r2), a["cierre"],
                    f"MK-ORI-{l['id']}", f"MK-LOC-{l['id']}", cc.get("codlocal", ""), cc.get("nombre_padron", ""), cc.get("direccion", ""), cc.get("lat", ""), cc.get("lon", ""), cc.get("confianza", "sin coordenada"),
                    sum(1 for x in INC if x["l"] == l["id"]), sum(1 for b in DEN if b["u"] == l["ubigeo_inei"]), sum(1 for a_ in JNE if a_["u"] == l["ubigeo_inei"]),
                    "Sí" if (RISK.get(l["ubigeo_inei"]) or {}).get("verificacion") else "No", rs["total"]])
    return out


def consultas_envios() -> list[list]:
    out = []
    por_k = defaultdict(list)
    for r in REG:
        if r["ts"] >= INICIO and r["corte"] in (2, 3) and r.get("consultas"):
            por_k[(r["k"], r["corte"])].append(r)
    cortes_de = defaultdict(set)
    for (k, cc) in por_k:
        cortes_de[k].add(cc)
    for r in REG:
        if not (r["ts"] >= INICIO and r["corte"] in (2, 3) and r.get("consultas")):
            continue
        g = por_k[(r["k"], r["corte"])]
        m = MEJOR[(r["k"], r["corte"])]
        usa = m is r and r["corte"] == 3
        tot = sum(r["consultas"].values())
        motivo = ("Único envío del orientador en el corte 3" if len(g) == 1 else "Es el de mayor total entre sus reenvíos del corte 3") if usa else (
            "No cuenta: el criterio usa solo el corte 3 (cierre, consultas de toda la jornada)" if r["corte"] == 2 else f"Reenvío: se usó el de mayor total (fila {m['fila']}, {sum(m['consultas'].values())} consultas)")
        out.append([r["fila"], f"{r['ts']:%d/%m/%Y %H:%M:%S}", titulo(NOMBRE.get(r["k"], r["k"])), ", ".join(lids_de(r)), f"Corte {r['corte']}", *[r["consultas"].get(t, 0) for t in TIPOS_CONSULTA], tot,
                    "Sí" if usa else "No", motivo, len(g), "Sí" if len(cortes_de[r["k"]]) > 1 else "No"])
    return out


# ---------------------------------------------------------------- 05 distritos
def tramo(n: int) -> str:
    return "—" if n <= 0 else TRAMOS[sum(1 for u in UMB if n >= u)]


def distritos() -> list[list]:
    us = set(RISK) | {x["u"] for x in INC} | {b["u"] for b in DEN} | {a["u"] for a in JNE} | {l["ubigeo_inei"] for l in ORIE["registros"]}
    out = []
    for u in sorted(us):
        d = RISK.get(u) or {}
        rs = rest_de(u)
        dep, prov, dist = ubic(u)
        ver = d.get("verificacion")
        gr = [x["id"] for x in INC if x["u"] == u and grave(x)]
        out.append([u, dep, prov, dist, f"POL-{u}", rs["total"], rs["reniec"], rs["jne"], tramo(rs["total"]), "Sí" if rs["total"] > 0 else "No", f"{ver['tipo']}: {fn(ver['domicilios'])} domicilios" if ver and ver.get("domicilios") is not None else ("Sí" if ver else "No"),
                    "Sí" if ver else "No", ", ".join(f"REA-{a['item']}" for a in d.get("alertas_rea", [])) or "—", ", ".join(a["id"] for a in JNE if a["u"] == u) or "—", ", ".join(l["id"] for l in ORIE["registros"] if l["ubigeo_inei"] == u) or "—",
                    len([x for x in INC if x["u"] == u]), len(gr), f"MK-BR-{u}" if any(b["u"] == u for b in DEN) else "—", f"MK-BV-{u}" if gr else "—"])
    return out


def alertas_previas() -> list[list]:
    out = []
    for b in BANJ["registros"]:
        out.append([b["id"], "Denuncia informada al equipo (REA)" if b["id"].startswith("REA-") else "Caso de la matriz con bandera roja", b["fecha_recepcion"], b["departamento"], b["provincia"], b["distrito"], b["ubigeo_inei"],
                    b.get("canal") or "", b.get("entidad_emisora") or "", b.get("categoria_rea") or "", "Sí" if b["conflictividad"] else "No", b.get("motivo_conflictividad") or "", b["resumen_publicable"],
                    f"MK-BR-{b['ubigeo_inei']}" if b["conflictividad"] else "No se dibuja", b.get("estado")])
    ids = {b["id"] for b in BANJ["registros"]}
    for d in RIESGO["distritos"]:
        for a in d["alertas_rea"]:
            if f"REA-{a['item']}" not in ids:
                out.append([f"REA-{a['item']}", "REA precargado (casos.json, corte 30/09/2026)", "—", d["departamento"], d["provincia"], d["distrito"], d["ubigeo_inei"], "Precargado REA", "", "", "Sin dato", a["motivo"], a["motivo"],
                            f"MK-BR-{d['ubigeo_inei']}", "Precargado"])
    return out


def alertas_jne() -> list[list]:
    return [[a["id"], a["departamento"], a["provincia"], a["distrito"], a["ubigeo_inei"], a["jee"], a["tipo_riesgo"], a["riesgo_violencia"], a["resolucion_territorial"], f"MK-{a['id']}"] for a in JNE]


# ---------------------------------------------------------------- KPIs (misma cuenta del dashboard)
def kpis() -> list[list]:
    loc = ORIE["registros"]
    total = sum(l["n_orientadores"] for l in loc)
    E = []
    for l in loc:
        a = J["avance"].get(l["id"], {"llegaron": 0, "cierre": 0})
        E.append([2 if i < a["cierre"] else 1 if i < a["llegaron"] else 0 for i in range(l["n_orientadores"])])
    ori = sum(1 for e in E for x in e if x >= 1)
    llego = sum(1 for e in E if all(x >= 1 for x in e))
    cerro = sum(1 for e in E if all(x == 2 for x in e))
    graves = sum(1 for x in INC if grave(x))
    c2, cs = J["corte2"], J["consultas"]
    return [
        [f"{ori} de {total}", "orientadores con llegada reportada", f"{llego} de {len(loc)} locales completos", "Orientadores del Forms con envío de llegada (corte 1) sobre los puestos de la base VF.", "03_FORMS_ORIENTADORES → hoja ORIENTADORES / LOCALES", "data/jornada.json → avance"],
        [fn(cs["total"]), "personas consultaron a los orientadores", "estimado reportado por los orientadores", "Consultas del corte 3 (cierre, total de toda la jornada); si un orientador reenvió el cierre, vale el de mayor total.", "04_CONSULTAS", "data/jornada.json → consultas"],
        [f"{c2['orientadores']} de {J['orientadores_base']}", "orientadores con registro del corte 2 (mediodía)", f"{c2['locales']} de {len(loc)} locales · {c2['distritos']} distritos", "Orientadores distintos con un envío del corte 2.", "03_FORMS_ORIENTADORES → hoja ENVIOS", "data/jornada.json → corte2"],
        [f"{cerro} de {len(loc)}", "locales con cierre enviado", "cierre hacia las 16:30", "Locales donde todos sus orientadores enviaron el corte 3 con hora de término.", "03_FORMS_ORIENTADORES → hoja LOCALES", "data/jornada.json → avance[].cierre"],
        [len(INC), "incidencias reportadas el domingo", f"{graves} graves · {len(INC) - graves} de atención", "Incidencias de jornada.json: las del Forms validadas y las de la matriz. Grave = tipo A / gravedad Alta / texto de violencia.", "01_INCIDENCIAS_CENTRAL → hoja CENTRAL (tipo «Incidencia»)", "data/jornada.json → incidencias"],
        [len(JNE), "alertas del JNE", f"riesgo de violencia · {len({a['u'] for a in JNE})} distritos · previas", "Alertas del JNE con riesgo de violencia.", "05_DISTRITOS_Y_ALERTAS → hoja ALERTAS_JNE", f"data/{ARCH['JNE_']}"],
        [len(DEN), "alertas RENIEC", f"denuncias en {len({b['u'] for b in DEN})} distritos · previas", "Banderas con riesgo de conflicto (REA informadas y de la matriz) + alertas REA precargadas que no están en ese archivo.", "01_INCIDENCIAS_CENTRAL → CENTRAL (tipo «Alerta RENIEC»)", f"data/{ARCH['BAN_']} + precarga/riesgo_previo.json"],
    ]


# ---------------------------------------------------------------- armado
def main() -> None:
    ahora = dt.datetime.now()
    carpeta = R / "seguimiento" / f"AUDITORIA_DASHBOARD_Q{ahora:%Y%m%d_%H%M}"
    carpeta.mkdir(parents=True, exist_ok=True)
    CEN, NOP, MVD = central(), no_publicadas(), matriz_vs_dashboard()
    anch_cen = {"TEXTO PUBLICADO EN EL DASHBOARD": 60, "TEXTO ORIGINAL EN LA FUENTE": 60, "CÓMO LLEGÓ AL DASHBOARD": 50, "DÓNDE SE VE (tablas, mapa, tarjetas)": 50, "REVISIÓN HUMANA": 45, "¿POR QUÉ ES GRAVE?": 38,
                "DÓNDE ESTÁ EL DATO ORIGINAL": 40, "FUENTE DE ORIGEN": 38, "OBSERVACIONES Y AVISOS": 45, "ARCHIVO Y CAMPO QUE LEE EL DASHBOARD": 38, "CAMBIOS ENTRE EL ORIGINAL Y LO PUBLICADO": 38}
    corte_txt = J["corte"]

    # ---- 01
    guarda(carpeta, "01_INCIDENCIAS_CENTRAL.xlsx", [
        leeme([("Para qué sirve", "Aquí está CADA incidencia y alerta que el dashboard muestra en su tabla «Detalle de las alertas e incidencias» o en el mapa, en el mismo orden de la tabla, con su fuente, su revisión humana y su marcador."),
               ("Cómo leerlo", "Una fila = una fila de la tabla del dashboard. Las columnas «ID DEL MARCADOR» y «¿SE DIBUJA EN EL MAPA?» dicen qué símbolo del mapa la representa (ver 02_MARCADORES_MAPA.xlsx). Las columnas de fuente dicen de dónde salió y cómo se cambió el texto."),
               ("Hojas", "CENTRAL: todo lo que se muestra · NO_PUBLICADAS: lo que existe en las fuentes pero NO está en el dashboard, con el motivo · MATRIZ_VS_DASHBOARD: cada caso de la matriz del equipo y dónde terminó."),
               ("Tipos", "Incidencia (domingo 4) = reportada hoy por orientadores o por la matriz · Alerta RENIEC = denuncia previa o caso con bandera roja · Alerta del JNE = riesgo reportado por los JEE antes de la jornada."),
               ("Mapa", "Solo se dibujan: banderas verdes (incidencias GRAVES del domingo, una por distrito), banderas rojas (alertas RENIEC, una por distrito) y el logo del JNE (una por alerta). Las incidencias de atención no se dibujan pero sí cuentan en cifras, gráficos y tablas."),
               ("Grave", "Una incidencia es grave si el equipo la etiquetó como grave (tipo A en el Forms / gravedad Alta en la matriz) o si su texto menciona violencia (amenaza, agresión, enfrentamiento, incendio, etc.). La columna «¿POR QUÉ ES GRAVE?» da el motivo exacto."),
               ("Uso interno", "Contiene nombres de orientadores y enlaces de la matriz. No es un archivo público. Los DNI y teléfonos de los textos originales se taparon."),
               ("Corte de datos", f"Dashboard {HTML.stem} · último envío del Forms: {corte_txt} · generado el {ahora:%d/%m/%Y %H:%M}.")]),
        ("CENTRAL", COLS_CENTRAL, CEN, anch_cen),
        ("NO_PUBLICADAS", COLS_NOPUB, NOP, {"TEXTO ORIGINAL": 60, "POR QUÉ NO ESTÁ (o cómo está cubierto)": 50}),
        ("MATRIZ_VS_DASHBOARD", ["ID EN LA HOJA", "FILA EN LA HOJA", "DISTRITO", "PROVINCIA", "DEPARTAMENTO", "FECHA", "BANDERA ROJA RENIEC (hoja)", "DÓNDE TERMINÓ EN EL DASHBOARD", "ID(S) EN EL DASHBOARD",
                                "LO QUE DICE LA HOJA en «¿Se ha registrado en Dashboard?»", "LO QUE DEBERÍA DECIR", "¿COINCIDE?", "DETALLE", "VALIDACIÓN"], MVD, {"DETALLE": 60})])

    # ---- 02
    mk_filas = [[m["id"], m["capa"], m["simbolo"], m["pestana"], *ubic(m["u"]), m["u"], m["lon"], m["lat"], m["origen_pos"], m["n"], m["ids"], m["rotulo"], m["interruptor"], m["datos"], m["obs"]] for m in MK]
    cols_mk = ["ID DEL MARCADOR", "CAPA", "SÍMBOLO", "PESTAÑA", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "UBIGEO INEI", "LONGITUD", "LATITUD", "DE DÓNDE SALE LA POSICIÓN", "N.º DE REGISTROS QUE REPRESENTA",
               "IDs DE LOS REGISTROS", "RÓTULO (texto al pasar el cursor)", "INTERRUPTOR QUE LO APAGA", "ARCHIVO DE DATOS", "OBSERVACIONES"]
    resumen_mk = Counter(m["capa"] for m in MK)
    guarda(carpeta, "02_MARCADORES_MAPA.xlsx", [
        leeme([("Para qué sirve", "Lista cada símbolo (marcador) que se dibuja en los mapas del dashboard, con su posición y los registros que representa."),
               ("Cómo leerlo", "MARCADORES: una fila por símbolo. REGISTROS_POR_MARCADOR: una fila por cada registro que hay detrás de un símbolo (una bandera puede agrupar varios). Use el «ID DEL REGISTRO» para buscarlo en 01_INCIDENCIAS_CENTRAL.xlsx."),
               ("Posición", "Los pines y las banderas se dibujan en el centro del distrito: es aproximado. Solo la pestaña «Locales de orientación» usa la coordenada del colegio (padrón MINEDU)."),
               ("Resumen", " · ".join(f"{k}: {v}" for k, v in resumen_mk.items()))]),
        ("MARCADORES", cols_mk, mk_filas, {"IDs DE LOS REGISTROS": 50, "DE DÓNDE SALE LA POSICIÓN": 50}),
        ("REGISTROS_POR_MARCADOR", ["ID DEL MARCADOR", "TIPO DE REGISTRO", "ID DEL REGISTRO", "TEXTO"], [[r["marcador"], r["tipo"], r["id"], r["texto"]] for r in MKREG], {"TEXTO": 90})])

    # ---- 03
    guarda(carpeta, "03_FORMS_ORIENTADORES.xlsx", [
        leeme([("Para qué sirve", "Muestra qué hizo cada orientador y cómo cada envío del Google Forms alimenta el dashboard (avance, incidencias, consultas)."),
               ("Hojas", "ENVIOS: cada fila del Forms · ORIENTADORES: los 51 orientadores de la base VF con su monitor operativo y quién falta · LOCALES: los 30 locales con su avance y sus marcadores."),
               ("Reglas", "Llegada = envío del corte 1 con hora de llegada. Cierre = envío del corte 3 con hora de término. Las filas de antes del 04/10 son pruebas y no cuentan. Para consultas ver 04_CONSULTAS.xlsx."),
               ("Uso interno", "Lleva nombres de orientadores y de monitores operativos; no publicar.")]),
        ("ENVIOS", ["FILA EN EL FORMS", "FECHA Y HORA DE ENVÍO", "ORIENTADOR", "LOCAL(ES) QUE ATIENDE (id)", "DISTRITO", "LOCAL (como lo eligió)", "CORTE", "HORA DE LLEGADA", "HORA DE TÉRMINO", "TIPO DE ENVÍO", "QUÉ ALIMENTA EN EL DASHBOARD",
                    "¿REPORTÓ INCIDENCIA?", "TEXTO DE LA INCIDENCIA (original)", "ID DE LA INCIDENCIA EN EL DASHBOARD", "ESTADO DE LA INCIDENCIA", "TOTAL DE CONSULTAS DEL ENVÍO", "¿CUENTA EN EL TOTAL DE CONSULTAS?"], envios(), {"TEXTO DE LA INCIDENCIA (original)": 60}),
        ("ORIENTADORES", ["ORIENTADOR", "MONITOR OPERATIVO", "OFICINA REGIONAL", "LOCAL(ES) (id)", "LOCAL(ES)", "DISTRITO", "DEPARTAMENTO", "LLEGADA (corte 1)", "REGISTRO DEL CORTE 2", "CIERRE (corte 3)", "ENVÍOS REALES", "CONSULTAS QUE CUENTAN"], orientadores()),
        ("LOCALES", ["ID DEL LOCAL", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "UBIGEO INEI", "LOCAL DE VOTACIÓN", "ORIENTADORES ASIGNADOS", "CON LLEGADA", "LLEGADA (X/N como en la tabla)", "CON REGISTRO DEL CORTE 2", "CON CIERRE",
                     "MARCADOR EN EL MAPA PRINCIPAL", "MARCADOR EN EL MAPA DE LOCALES", "CÓDIGO DE LOCAL (MINEDU)", "NOMBRE EN EL PADRÓN MINEDU", "DIRECCIÓN", "LATITUD", "LONGITUD", "CONFIANZA DE LA COORDENADA", "INCIDENCIAS DEL DOMINGO EN EL LOCAL",
                     "ALERTAS RENIEC EN EL DISTRITO", "ALERTAS JNE EN EL DISTRITO", "¿DISTRITO VERIFICADO?", "RESTITUIDOS EN EL DISTRITO"], locales_fila())])

    # ---- 04
    ce = consultas_envios()
    usados = [f for f in ce if f[11] == "Sí"]
    por_tipo = {t: [sum(f[5 + i] for f in usados if f[4] == cor) for cor in ("Corte 2", "Corte 3")] for i, t in enumerate(TIPOS_CONSULTA)}
    bruto = {t: sum(f[5 + i] for f in ce) for i, t in enumerate(TIPOS_CONSULTA)}
    total = sum(sum(v) for v in por_tipo.values())
    res_tipo = [[t, sum(v), f"{100 * sum(v) / total:.1f} %" if total else "", v[0], v[1], bruto[t], bruto[t] - sum(v)] for t, v in por_tipo.items()] + [["TOTAL", total, "100 %", sum(v[0] for v in por_tipo.values()), sum(v[1] for v in por_tipo.values()), sum(bruto.values()), sum(bruto.values()) - total]]
    por_local = defaultdict(lambda: [0, 0])
    for f in usados:
        for lid in str(f[3]).split(", "):
            por_local[lid][0 if f[4] == "Corte 2" else 1] += f[5 + 5] / max(1, len(str(f[3]).split(", ")))
    res_local = [[lid, LOCALES[lid]["distrito"], LOCALES[lid]["local"], round(v[0], 1), round(v[1], 1), round(sum(v), 1)] for lid, v in sorted(por_local.items())]
    guarda(carpeta, "04_CONSULTAS.xlsx", [
        leeme([("Para qué sirve", "Explica de dónde sale el total de «personas que consultaron» del dashboard y cómo se descartaron los duplicados."),
               ("Regla", "Criterio «solo corte 3» (decisión del 04/10/2026): cuentan las consultas del cierre, que el formulario pide como total de toda la jornada. Si un orientador envió el cierre varias veces, se usa el envío de mayor total. Los envíos del corte 2 no suman. La hoja ENVIOS_CON_CONSULTAS marca con Sí/No cada envío y da el motivo."),
               ("Atención", "Un orientador que solo reportó el corte 2 no suma. Algunos tienen en el corte 3 menos que en el 2 (columna «¿Envió también en el otro corte?»); con este criterio vale lo que declararon al cierre."),
               ("Hojas", "RESUMEN_POR_TIPO · RESUMEN_POR_LOCAL (consultas repartidas por local; si un orientador atiende dos locales se divide en partes iguales) · ENVIOS_CON_CONSULTAS."),
               ("Cuadre", f"Total del dashboard (data/jornada.json): {J['consultas']['total']} · total de este archivo: {total}.")]),
        ("RESUMEN_POR_TIPO", ["TIPO DE CONSULTA", "TOTAL QUE CUENTA", "% DEL TOTAL", "CORTE 2", "CORTE 3", "SUMA BRUTA (todos los envíos)", "DUPLICADOS DESCARTADOS"], res_tipo),
        ("RESUMEN_POR_LOCAL", ["ID DEL LOCAL", "DISTRITO", "LOCAL", "CORTE 2", "CORTE 3", "TOTAL"], res_local),
        ("ENVIOS_CON_CONSULTAS", ["FILA EN EL FORMS", "FECHA Y HORA", "ORIENTADOR", "LOCAL(ES)", "CORTE", *TIPOS_CONSULTA, "TOTAL DEL ENVÍO", "¿CUENTA?", "MOTIVO", "ENVÍOS DEL ORIENTADOR EN ESTE CORTE", "¿ENVIÓ TAMBIÉN EN EL OTRO CORTE?"], ce, {"MOTIVO": 55})])

    # ---- 05
    guarda(carpeta, "05_DISTRITOS_Y_ALERTAS_PREVIAS.xlsx", [
        leeme([("Para qué sirve", "Bases por distrito y alertas previas a la jornada que alimentan el mapa (colores de restituidos, contorno de verificación, banderas rojas y logos del JNE)."),
               ("Hojas", "DISTRITOS: una fila por distrito con datos (relleno de color por restituidos, contorno rojo por verificación, alertas y marcadores) · ALERTAS_RENIEC: denuncias REA y casos de la matriz con bandera roja · ALERTAS_JNE: riesgos que informó el JNE."),
               ("Colores", "El relleno del distrito sigue 4 tramos de restituidos: 1–99, 100–499, 500–2 999 y ≥ 3 000. El contorno rojo marca los distritos donde RENIEC hizo verificación domiciliaria."),
               ("Fuentes", f"Restituidos: {RIESGO['fuentes']['restituidos']} · verificaciones: {RIESGO['fuentes']['verificaciones']} · alertas REA: {RIESGO['fuentes']['denuncias_con_alerta']}.")]),
        ("DISTRITOS", ["UBIGEO INEI", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "ID DEL POLÍGONO EN EL MAPA", "RESTITUIDOS (total)", "RESTITUIDOS RENIEC", "RESTITUIDOS JNE", "TRAMO DE COLOR", "¿SE PINTA RELLENO?", "VERIFICACIÓN DOMICILIARIA", "¿CONTORNO ROJO?",
                       "ALERTAS REA PREVIAS", "ALERTAS JNE", "LOCALES DE ORIENTACIÓN", "INCIDENCIAS DEL DOMINGO", "DE ELLAS, GRAVES", "BANDERA ROJA (marcador)", "BANDERA VERDE (marcador)"], distritos()),
        ("ALERTAS_RENIEC", ["ID", "ORIGEN", "FECHA DE RECEPCIÓN", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "UBIGEO INEI", "CANAL", "ENTIDAD EMISORA", "CATEGORÍA REA", "¿RIESGO DE CONFLICTO?", "MOTIVO", "TEXTO PUBLICADO", "MARCADOR EN EL MAPA", "ESTADO"],
         alertas_previas(), {"TEXTO PUBLICADO": 70, "MOTIVO": 45}),
        ("ALERTAS_JNE", ["ID", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "UBIGEO INEI", "JEE", "TIPO DE RIESGO", "¿RIESGO DE VIOLENCIA?", "RESOLUCIÓN DEL TERRITORIO", "MARCADOR EN EL MAPA"], alertas_jne(), {"TIPO DE RIESGO": 60})])

    # ---- 00 (guia, KPIs, controles)
    graves = [x for x in INC if grave(x)]
    control = [
        ("Filas de la tabla del dashboard", len(TABLA), len(INC) + len(DEN) + len(JNE), "incidencias + alertas RENIEC + alertas JNE"),
        ("Filas de CENTRAL", len(CEN), len(TABLA), "una por fila de la tabla"),
        ("Incidencias en jornada.json", len(INC), len([f for f in CEN if f[2].startswith("Incidencia")]), "CENTRAL tipo Incidencia"),
        ("Banderas verdes (distritos con grave)", len({x["u"] for x in graves}), len([m for m in MK if m["id"].startswith("MK-BV")]), "marcadores MK-BV"),
        ("Banderas rojas (distritos con alerta)", len({b["u"] for b in DEN}), len([m for m in MK if m["id"].startswith("MK-BR")]), "marcadores MK-BR"),
        ("Logos JNE", len(JNE), len([m for m in MK if m["id"].startswith("MK-JNE")]), "marcadores MK-JNE"),
        ("Pines de orientación", len(ORIE["registros"]), len([m for m in MK if m["id"].startswith("MK-ORI")]), "marcadores MK-ORI"),
        ("Consultas (dashboard vs. este análisis)", J["consultas"]["total"], total, "04_CONSULTAS"),
        ("Casos de la matriz clasificados", len(CASOS), len(MVD), "MATRIZ_VS_DASHBOARD"),
        ("Registros detrás de los marcadores", sum(m["n"] for m in MK), len(MKREG), "REGISTROS_POR_MARCADOR"),
    ]
    guarda(carpeta, "00_GUIA_Y_KPIS_DEL_DASHBOARD.xlsx", [
        leeme([("Qué es esto", f"Paquete de auditoría del dashboard https://monitoreo-erm-2026.vercel.app (versión {HTML.stem}), con datos al {corte_txt}. Cada archivo explica una parte; este resume cómo se construye todo."),
               ("Archivos", "01_INCIDENCIAS_CENTRAL: todas las incidencias y alertas con fuente y marcador · 02_MARCADORES_MAPA: cada símbolo del mapa · 03_FORMS_ORIENTADORES: envíos, orientadores y locales · 04_CONSULTAS: total de personas orientadas y duplicados · 05_DISTRITOS_Y_ALERTAS_PREVIAS: capas de distrito y alertas previas."),
               ("Ruta de los datos", "1) Los orientadores llenan el Google Forms (llegada, consultas, incidencias, cierre). 2) El equipo llena la matriz de denuncias (Google Sheet «Casos») y revisa las incidencias del Forms. 3) Un script cruza todo con la base de orientadores (VF), "
                                     "los locales y los distritos. 4) Se generan los archivos de datos (data/jornada.json, banderas_rea, etc.). 5) El dashboard los lee y dibuja el mapa, las tarjetas, los gráficos y las tablas."),
               ("Qué se dibuja en el mapa", "Pines de orientación (uno por local), banderas rojas (alertas RENIEC previas y casos con bandera roja), logos JNE (alertas del JNE) y banderas verdes (solo incidencias graves del domingo). Más dos capas de distrito: relleno por restituidos y contorno rojo por verificación."),
               ("Qué es grave", "Etiqueta grave del equipo (tipo A en el Forms, gravedad Alta en la matriz) o texto que menciona violencia. Solo las graves llevan bandera verde."),
               ("Consultas", "Solo corte 3 (cierre, consultas de toda la jornada); ante reenvíos vale el de mayor total."),
               ("Orientadores y locales", "Base: MONITOREO_REPORTE 041026_VF (51 orientadores en 30 locales; Erick Flores atiende dos locales)."),
               ("Qué NO se publica", "Nombres de personas, DNI, teléfonos y enlaces. Esos datos solo están en estos archivos internos y en las fuentes."),
               ("Pendientes conocidos", "Casos de la matriz sin decisión (ver NO_PUBLICADAS), IDs repetidos en la hoja (C25 y C26) y 6 locales con coordenada «por revisar».")]),
        ("KPIS", ["CIFRA QUE SE VE", "TARJETA", "SUBTÍTULO", "CÓMO SE CALCULA", "DÓNDE VERLO EN ESTE PAQUETE", "ARCHIVO DE DATOS"], kpis(), {"CÓMO SE CALCULA": 70}),
        ("CONTROLES", ["CONTROL", "VALOR EN LOS DATOS", "VALOR EN EL PAQUETE", "RESULTADO", "COMENTARIO"], [[a, b, cc, "OK" if b == cc else "REVISAR", d] for a, b, cc, d in control]),
        ("FUENTES_Y_ARCHIVOS", ["FUENTE", "ARCHIVO", "QUÉ ALIMENTA", "ÚLTIMA MODIFICACIÓN"], [
            ["Google Forms de orientadores", str(FORMS.relative_to(R)), "avance, consultas, incidencias del Forms", f"{dt.datetime.fromtimestamp(FORMS.stat().st_mtime):%d/%m/%Y %H:%M}"],
            ["Matriz de denuncias (Google Sheet «Casos»)", str(MATRIZ.relative_to(R)), "incidencias de la matriz y banderas rojas MAT-", f"{dt.datetime.fromtimestamp(MATRIZ.stat().st_mtime):%d/%m/%Y %H:%M}"],
            ["Revisión humana del Forms", "data/entrada/incidencias_forms.csv", "qué incidencias del Forms se publican, su tipo y su resumen", ""],
            ["Revisión humana de la matriz", "data/entrada/validaciones_matriz.csv", "qué casos de la matriz se publican, canal, motivo y texto", ""],
            ["Base de orientadores VF", str(BASE.relative_to(R.parent)), "51 orientadores, locales y monitores", f"{dt.datetime.fromtimestamp(BASE.stat().st_mtime):%d/%m/%Y %H:%M}"],
            ["Datos de la jornada", f"data/{ARCH['JOR']}", "avance, consultas e incidencias del domingo", ""],
            ["Banderas rojas", f"data/{ARCH['BAN_']}", "alertas RENIEC (REA informadas y casos de la matriz)", ""],
            ["Alertas del JNE", f"data/{ARCH['JNE_']}", "logos JNE", ""],
            ["Locales de orientación", f"data/{ARCH['ORIE']} y data/{ARCH['COORD']}", "pines, locales y coordenadas de los colegios", ""],
            ["Precarga de riesgos", "dashboard/public/data/precarga/riesgo_previo.json", "restituidos, verificación y alertas REA previas por distrito", ""],
            ["Mapa de distritos", "dashboard/public/data/precarga/distritos_pais.geojson", "polígonos y centros de distrito", ""]])])

    print(f"carpeta: {carpeta}")
    for nombre, b, cc, d in control:
        print(f"  {'OK ' if b == cc else 'REVISAR'} {nombre}: {b} / {cc}")
    print(f"marcadores: {dict(resumen_mk)} | incidencias {len(INC)} ({len(graves)} graves) | alertas RENIEC {len(DEN)} | JNE {len(JNE)} | no publicadas {len(NOP)}")


if __name__ == "__main__":
    main()
