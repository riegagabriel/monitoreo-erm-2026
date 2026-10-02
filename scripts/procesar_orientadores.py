"""Despliegue de orientacion: de ORIENTADORES_AL_1_10.xlsx a datos publicables y a la lista privada del formulario.

Salidas:
  data/orientacion_locales.json            PUBLICO: un registro por local (sin nombres, DNI, celular ni correo)
  formulario/orientadores_para_pegar.csv   PRIVADO (ignorado por Git): etiqueta con el nombre para el desplegable

Territorio: tripleta departamento-provincia-distrito contra el catalogo INEI. Un caso que no empareje se resuelve
SOLO con data/equivalencias_orientadores.csv (decision explicita, nunca por coincidencia parcial); si no, aborta.
"""
from __future__ import annotations

import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "ORIENTADORES_AL_1_10.xlsx"
EQUIV = RAIZ / "data" / "equivalencias_orientadores.csv"
LOOKUP = RAIZ.parent / "DENUNCIAS_REA" / "data" / "geodata" / "distritos_lookup.parquet"
PUBLICO = RAIZ / "data" / "orientacion_locales.json"
PRIVADO = RAIZ / "formulario" / "orientadores_para_pegar.csv"
CORTE = "01/10/2026"  # nombre de la hoja de la fuente: «AL 01.10.26»
COLUMNAS = ["OFICINA REGIONAL", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "NOMBRE DEL LOCAL", "NOMBRES Y APELLIDOS"]


def norm(x) -> str:
    s = unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", s)).strip()


def leer(ruta: Path = FUENTE) -> list[dict]:
    ws = openpyxl.load_workbook(ruta, data_only=True).active
    cab = [str(c).strip() if c else None for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    faltan = [c for c in COLUMNAS if c not in cab]
    if faltan:
        raise SystemExit(f"Faltan columnas en la lista de orientadores: {faltan}")
    filas = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(cab, r))
        if d.get("NOMBRES Y APELLIDOS"):
            filas.append({k: str(d[k]).strip() if d.get(k) is not None else "" for k in COLUMNAS})
    return filas


def leer_equivalencias(ruta: Path = EQUIV) -> dict[tuple, tuple]:
    if not ruta.exists():
        return {}
    with ruta.open(encoding="utf-8-sig", newline="") as f:
        return {(norm(r["dep_origen"]), norm(r["prov_origen"]), norm(r["dist_origen"])): (r["dep"], r["prov"], r["dist"])
                for r in csv.DictReader(f)}


def construir(filas: list[dict], catalogo: pd.DataFrame, equiv: dict[tuple, tuple]) -> tuple[list[dict], list[dict]]:
    idx = {(norm(r.DEPARTAMEN), norm(r.PROVINCIA), norm(r.DISTRITO)): r for r in catalogo.itertuples()}
    locales: dict[tuple, dict] = {}
    privado = []
    for f in filas:
        clave = (norm(f["DEPARTAMENTO"]), norm(f["PROVINCIA"]), norm(f["DISTRITO"]))
        if clave in equiv:
            clave = tuple(norm(x) for x in equiv[clave])
        c = idx.get(clave)
        if c is None:
            raise SystemExit(f"No se puede ubicar {f['DEPARTAMENTO']!r} / {f['PROVINCIA']!r} / {f['DISTRITO']!r}: "
                             f"agregue una equivalencia explicita en {EQUIV.name}")
        k = (str(c.UBIGEO_INEI), norm(f["NOMBRE DEL LOCAL"]))
        loc = locales.setdefault(k, {
            "ubigeo_inei": str(c.UBIGEO_INEI), "departamento": str(c.DEPARTAMEN), "provincia": str(c.PROVINCIA),
            "distrito": str(c.DISTRITO), "local": f["NOMBRE DEL LOCAL"].strip(), "oficina_regional": f["OFICINA REGIONAL"].strip(),
            "lon": round(float(c.lon), 5), "lat": round(float(c.lat), 5), "n_orientadores": 0,
        })
        loc["n_orientadores"] += 1
        privado.append({"ubigeo_inei": loc["ubigeo_inei"], "local": loc["local"], "nombre": f["NOMBRES Y APELLIDOS"].strip()})
    publico = sorted(locales.values(), key=lambda l: (l["departamento"], l["provincia"], l["distrito"], l["local"]))
    for i, l in enumerate(publico, 1):
        l["id"] = f"L-{i:02d}"
    por_local = {(l["ubigeo_inei"], l["local"]): l["id"] for l in publico}
    for p in privado:
        p["local_id"] = por_local[(p["ubigeo_inei"], p["local"])]
    privado.sort(key=lambda p: (p["local_id"], p["nombre"]))
    return publico, privado


def main() -> None:
    filas = leer()
    cat = pd.read_parquet(LOOKUP, columns=["UBIGEO_INEI", "DEPARTAMEN", "PROVINCIA", "DISTRITO", "lon", "lat"])
    publico, privado = construir(filas, cat, leer_equivalencias())
    por_dist = defaultdict(list)
    for l in publico:
        por_dist[l["ubigeo_inei"]].append(l["id"])
    PUBLICO.parent.mkdir(parents=True, exist_ok=True)
    PUBLICO.write_text(json.dumps({
        "fuente": "Lista de orientadores desplegados por RENIEC", "corte": CORTE,
        "orientadores": len(filas), "locales": len(publico), "distritos": len(por_dist), "registros": publico,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    PRIVADO.parent.mkdir(parents=True, exist_ok=True)
    with PRIVADO.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["etiqueta", "local_id", "ubigeo_inei"])
        for l in publico:
            for p in (x for x in privado if x["local_id"] == l["id"]):
                w.writerow([f"{l['distrito']} · {l['local']} · {p['nombre']}", l["id"], l["ubigeo_inei"]])
    print(f"{len(filas)} orientadores, {len(publico)} locales, {len(por_dist)} distritos -> {PUBLICO.name}; lista privada -> {PRIVADO.name}")


if __name__ == "__main__":
    sys.exit(main())
