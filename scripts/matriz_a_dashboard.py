"""Matriz de casos (hoja de Google «Casos») -> registros para el dashboard. Sin --aplicar es una SIMULACION: no toca el dashboard ni los JSON publicados.

Uso: python scripts/matriz_a_dashboard.py [matriz.xlsx] [--descargar] [--aplicar]
Entrada por defecto: data/entrada/matriz_casos.xlsx (ignorada por Git). Con --descargar la baja del enlace guardado en data/entrada/matriz_url.txt.
Salida: reporte en pantalla y data/entrada/matriz_simulacion.json (lo que se publicaria, caso por caso).
Con --aplicar, solo los casos «LISTO PARA PUBLICAR»: verdes -> data/entrada/incidencias_matriz.json (lo lee corte.py); rojas -> se agregan a
data/banderas_rea_Q20261002.json (los ids MAT-* anteriores se reemplazan, los REA-* no se tocan).
data/entrada/validaciones_matriz.csv (id, validado, motivo, canal, texto) completa las columnas que la hoja aun no tiene; el valor de la hoja manda
si existe, salvo «texto», que reemplaza a la descripcion de la hoja.

Reglas:
- Destino por «Fecha reporte»: el dia de la jornada (04/10/2026) -> bandera VERDE (incidencia del domingo, formato de jornada.json);
  fecha anterior con «Bandera roja RENIEC = Si» -> bandera ROJA (alerta previa, formato de banderas_rea.json); lo demas no se publica.
- Solo se publica lo que el equipo marca «SI» en la columna «Validado para dashboard». Lo demas queda como pendiente.
- Ubigeo RENIEC -> INEI por la hoja Ref_Distritos de la misma matriz; el ubigeo debe coincidir con el nombre del distrito.
- Texto: solo «Descripcion resumida» (maximo 280 caracteres en verde, 700 en roja); sin DNI, telefonos, enlaces ni nombres.
- Verde exige «Canal dashboard» (ori/rea/ent/pre/tel) y gravedad Alta/Media/Baja; roja exige «Motivo (1 frase)» (maximo 160).
- Un distrito que ya tiene bandera roja publicada se avisa como posible duplicado (la decision es del equipo).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "data" / "entrada" / "matriz_casos.xlsx"
URL = RAIZ / "data" / "entrada" / "matriz_url.txt"
SIMULACION = RAIZ / "data" / "entrada" / "matriz_simulacion.json"
VALIDACIONES = RAIZ / "data" / "entrada" / "validaciones_matriz.csv"
INC_MATRIZ = RAIZ / "data" / "entrada" / "incidencias_matriz.json"
GEOJSON = RAIZ / "dashboard" / "public" / "data" / "precarga" / "distritos_pais.geojson"
BANDERAS = RAIZ / "data" / "banderas_rea_Q20261002.json"
LOCALES = RAIZ / "data" / "orientacion_locales.json"
PROHIBIDAS = RAIZ.parent / "DENUNCIAS_REA" / "config" / "frases_prohibidas.local.txt"
DIA = dt.date(2026, 10, 4)
MAX_VERDE, MAX_ROJA, MAX_MOTIVO = 280, 700, 160
CANALES = {"ori": "Orientadores", "rea": "Denuncias REA de hoy", "ent": "Entidades del sistema", "pre": "Prensa y redes sociales", "tel": "Llamadas o mesa de ayuda"}
GRAVEDAD = {"ALTA": 2, "MEDIA": 1, "BAJA": 1}
COLS = {"id": "ID", "dep": "DEPARTAMENTO", "prov": "PROVINCIA", "dist": "DISTRITO", "ubigeo": "UBIGEO RENIEC", "local": "LOCAL DE VOTACION",
        "fecha": "FECHA REPORTE", "hora": "HORA", "desc": "DESCRIPCION RESUMIDA", "gravedad": "GRAVEDAD", "bandera": "BANDERA ROJA",
        "remitente": "REMITENTE", "responsable": "RESPONSABLE", "validado": "VALIDADO PARA DASHBOARD", "motivo": "MOTIVO", "canal": "CANAL DASHBOARD"}
OBLIGATORIAS = ("id", "dist", "fecha", "desc")


def nz(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", s)).strip()


def leer_casos(ws) -> tuple[list[dict], list[str]]:
    filas = list(ws.iter_rows(values_only=True))
    cab = [nz(c) for c in filas[0]]
    ix, faltan = {}, []
    for k, prefijo in COLS.items():
        i = next((j for j, h in enumerate(cab) if h == prefijo or (h and h.startswith(prefijo))), None)
        ix[k] = i
        if i is None:
            faltan.append(k)
    if any(k in faltan for k in OBLIGATORIAS):
        raise SystemExit(f"Faltan columnas obligatorias en la hoja Casos: {[COLS[k] for k in OBLIGATORIAS if k in faltan]}")
    out = []
    for r in filas[1:]:
        if r[ix["id"]] in (None, ""):
            continue
        out.append({k: (r[i] if i is not None else None) for k, i in ix.items()})
    return out, [COLS[k] for k in faltan]


def aplicar_validaciones(casos: list[dict], ruta: Path) -> None:
    if not ruta.exists():
        return
    with ruta.open(encoding="utf-8", newline="") as f:
        ov = {x["id"].strip(): x for x in csv.DictReader(f)}
    ids = {str(c["id"]).strip() for c in casos}
    if set(ov) - ids:
        raise SystemExit(f"validaciones_matriz.csv trae ids que no estan en la matriz: {sorted(set(ov) - ids)}")
    for c in casos:
        x = ov.get(str(c["id"]).strip())
        if not x:
            continue
        for k in ("validado", "motivo", "canal"):
            if c.get(k) in (None, "") and x.get(k, "").strip():
                c[k] = x[k].strip()
        if x.get("texto", "").strip():
            c["desc"] = x["texto"].strip()


def cargar_ref(wb) -> tuple[dict, dict]:
    """RENIEC -> INEI y (dep, prov, dist) -> INEI, desde la hoja Ref_Distritos de la matriz."""
    ws = next(w for w in wb.worksheets if w.title.startswith("Ref_Distritos"))
    filas = list(ws.iter_rows(values_only=True))
    cab = [nz(c) for c in filas[0]]
    i_inei = next(j for j, h in enumerate(cab) if h.startswith("UBIGEO INEI"))
    r2i, nombres = {}, {}
    for r in filas[1:]:
        if r[0] and r[i_inei]:
            r2i[str(r[0])] = str(r[i_inei])
            nombres[(nz(r[1]), nz(r[2]), nz(r[3]))] = str(r[i_inei])
    return r2i, nombres


def parse_fecha(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", str(v or "").strip()) or re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", str(v or "").strip())
    if not m:
        return None
    a, b, c = (int(x) for x in m.groups())
    return dt.date(a, b, c) if a > 31 else dt.date(c, b, a)


def parse_hora(v) -> float | None:
    if isinstance(v, (dt.datetime, dt.time)):
        return round(v.hour + v.minute / 60, 2)
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", str(v or "").strip())
    return round(int(m.group(1)) + int(m.group(2)) / 60, 2) if m and int(m.group(1)) < 24 else None


def nombres_internos(casos: list[dict]) -> set[str]:
    """Palabras con mayuscula inicial de las columnas Remitente y Responsable (personas del equipo) mas la lista local de frases prohibidas."""
    out = set()
    for c in casos:
        for campo in ("remitente", "responsable"):
            out |= {nz(t) for t in re.findall(r"[A-ZÁÉÍÓÚÑ][a-záéíóúñ]{3,}", str(c.get(campo) or ""))}
    if PROHIBIDAS.exists():
        out |= {nz(l) for l in PROHIBIDAS.read_text(encoding="utf-8").splitlines() if l.strip() and not l.startswith("#")}
    return out - {""}


def revisar_texto(t: str, prohibidos: set[str]) -> list[str]:
    p = []
    if re.search(r"\d{8,}", t.replace(" ", "")):
        p.append("tiene una cifra de 8 o mas digitos (posible DNI o telefono)")
    if re.search(r"https?://|www\.|@\w", t):
        p.append("tiene un enlace o usuario de red social")
    plano = f" {nz(t)} "
    hallados = sorted(n for n in prohibidos if f" {n} " in plano)
    if hallados:
        p.append(f"menciona nombres que no deben publicarse: {hallados}")
    return p


def ubicar(c: dict, ctx: dict) -> tuple[list[str], list[str]]:
    """Devuelve (ubigeos INEI, errores). Cada distrito del campo (separados por «;») se busca por nombre; el ubigeo RENIEC debe coincidir."""
    nombres = [d.strip() for d in str(c.get("dist") or "").split(";") if d.strip()]
    if not nombres:
        return [], []
    ubigeos, errores = [], []
    for d in nombres:
        u = ctx["nombres"].get((nz(c.get("dep")), nz(c.get("prov")), nz(d)))
        if u is None:
            errores.append(f"distrito «{d}» no figura en Ref_Distritos con ese departamento y provincia")
        else:
            ubigeos.append(u)
    cod = str(c.get("ubigeo") or "").strip()
    if cod and not errores:
        if ctx["r2i"].get(cod) not in ubigeos:
            errores.append(f"el ubigeo RENIEC {cod} no corresponde al distrito escrito")
    return ubigeos, errores


def local_id(c: dict, u: str, ctx: dict) -> tuple[str | None, str | None]:
    nombre = nz(c.get("local"))
    if not nombre:
        return None, None
    for l in ctx["locales"]:
        if l["ubigeo_inei"] == u and (nombre in nz(l["local"]) or nz(l["local"]) in nombre):
            return l["id"], None
    return None, f"el local «{c.get('local')}» no esta entre los locales con orientadores de ese distrito"


def evaluar(c: dict, ctx: dict) -> dict:
    cid = str(c["id"]).strip()
    res = {"id": cid, "distrito": str(c.get("dist") or ""), "destino": None, "estado": None, "errores": [], "avisos": [], "registros": []}
    fecha, bandera = parse_fecha(c.get("fecha")), nz(c.get("bandera"))
    if fecha is None:
        res["errores"].append(f"«Fecha reporte» no es una fecha valida: {c.get('fecha')!r}")
    elif fecha == DIA:
        res["destino"] = "verde"
    elif bandera == "SI":
        res["destino"] = "roja"
    else:
        res["estado"] = "NO APLICA" if bandera == "NO" else "BANDERA POR DEFINIR"
        return res
    ubigeos, errs = ubicar(c, ctx)
    res["errores"] += errs
    if not ubigeos and not errs:
        if res["destino"] == "roja":
            res["errores"].append("sin distrito: la bandera roja necesita ubicacion")
        else:
            res["avisos"].append("sin distrito: cuenta en los totales pero no se dibuja")
    desc = str(c.get("desc") or "").strip()
    limite = MAX_VERDE if res["destino"] == "verde" else MAX_ROJA
    if not desc:
        res["errores"].append("falta la descripcion resumida")
    elif len(desc) > limite:
        res["errores"].append(f"la descripcion mide {len(desc)} caracteres y el maximo es {limite}")
    res["errores"] += [f"descripcion: {p}" for p in revisar_texto(desc, ctx["prohibidos"])]
    if res["destino"] == "verde":
        canal, g, h = nz(c.get("canal")).lower(), GRAVEDAD.get(nz(c.get("gravedad"))), parse_hora(c.get("hora"))
        if canal not in CANALES:
            res["errores"].append("falta «Canal dashboard» (ori, rea, ent, pre o tel)")
        if g is None:
            res["errores"].append("la gravedad debe ser Alta, Media o Baja")
        if h is None:
            res["errores"].append(f"«Hora» no es HH:mm: {c.get('hora')!r}")
        if not res["errores"]:
            for n, u in enumerate(ubigeos or [""], start=1):
                lid, aviso = local_id(c, u, ctx) if u else (None, None)
                if aviso:
                    res["avisos"].append(aviso)
                res["registros"].append({"id": f"MAT-{cid}" + (f"-{n}" if len(ubigeos) > 1 else ""), "h": h, "g": g, "c": canal, "u": u, "l": lid, "t": desc})
    else:
        motivo = str(c.get("motivo") or "").strip()
        if not motivo:
            res["errores"].append("falta «Motivo (1 frase)» de la bandera roja")
        elif len(motivo) > MAX_MOTIVO:
            res["errores"].append(f"el motivo mide {len(motivo)} caracteres y el maximo es {MAX_MOTIVO}")
        res["errores"] += [f"motivo: {p}" for p in revisar_texto(motivo, ctx["prohibidos"])]
        if not res["errores"]:
            canal = nz(c.get("canal")).lower()
            for n, u in enumerate(ubigeos, start=1):
                g = ctx["geo"][u]
                res["registros"].append({"id": f"MAT-{cid}" + (f"-{n}" if len(ubigeos) > 1 else ""), "item": None, "fecha_recepcion": fecha.strftime("%d/%m/%Y"),
                                         "fecha_asumida": False, "departamento": g["dep"], "provincia": g["prov"], "distrito": g["dist"], "ubigeo_inei": u,
                                         "lon": g["lon"], "lat": g["lat"], "categoria_rea": None, "conflictividad": True, "motivo_conflictividad": motivo,
                                         "fuente_dashboard": "conflictividad", "resumen_publicable": desc, "n_ciudadanos": None, "entidad_emisora": None,
                                         "canal": CANALES.get(canal, "Matriz de monitoreo"), "estado": "VALIDADO", "simulado": False})
    for u in ubigeos:
        if res["destino"] == "roja" and ctx["banderas"].get(u):
            res["avisos"].append(f"posible duplicado: el distrito ya tiene {', '.join(ctx['banderas'][u])}")
    validado = nz(c.get("validado")) == "SI"
    res["estado"] = "ERROR" if res["errores"] else "LISTO PARA PUBLICAR" if validado else "PENDIENTE DE VALIDAR"
    return res


def contexto(wb, casos: list[dict]) -> dict:
    r2i, nombres = cargar_ref(wb)
    geo = {f["properties"]["u"]: f["properties"] for f in json.loads(GEOJSON.read_text(encoding="utf-8"))["features"]}
    ban = json.loads(BANDERAS.read_text(encoding="utf-8"))["registros"]
    por_u: dict[str, list[str]] = {}
    for b in ban:
        if b.get("conflictividad") in (True, "True") and not str(b["id"]).startswith("MAT-"):
            por_u.setdefault(b["ubigeo_inei"], []).append(b["id"])
    return {"r2i": r2i, "nombres": nombres, "geo": geo, "banderas": por_u, "prohibidos": nombres_internos(casos),
            "locales": json.loads(LOCALES.read_text(encoding="utf-8"))["registros"]}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("ruta", nargs="?", type=Path, default=ENTRADA)
    ap.add_argument("--descargar", action="store_true")
    ap.add_argument("--aplicar", action="store_true")
    a = ap.parse_args()
    if a.descargar:
        if not URL.exists():
            raise SystemExit(f"Falta {URL} con el enlace de publicacion (…/pub?output=xlsx)")
        a.ruta.parent.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(URL.read_text(encoding="utf-8").strip(), a.ruta)
    wb = openpyxl.load_workbook(a.ruta, data_only=True)
    casos, sin_columna = leer_casos(wb["Casos"])
    aplicar_validaciones(casos, VALIDACIONES)
    ctx = contexto(wb, casos)
    res = [evaluar(c, ctx) for c in casos]
    SIMULACION.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(res)} casos en la matriz | {'APLICAR' if a.aplicar else 'SIMULACION: no se escribio nada en el dashboard'}")
    if sin_columna:
        print("Columnas que la hoja aun no tiene:", sin_columna)
    for r in res:
        notas = "; ".join(r["errores"] + r["avisos"])
        print(f"{r['id']:4} {r['distrito'][:26]:26} {r['destino'] or '-':6} {r['estado']:22} {notas}")
    cuenta = {}
    for r in res:
        cuenta[r["estado"]] = cuenta.get(r["estado"], 0) + 1
    listos = [r for r in res if r["estado"] == "LISTO PARA PUBLICAR"]
    print("Resumen:", cuenta, "| se publicarian:", sum(len(r["registros"]) for r in listos), "registros")
    if a.aplicar:
        verdes = [x for r in listos if r["destino"] == "verde" for x in r["registros"]]
        rojas = [x for r in listos if r["destino"] == "roja" for x in r["registros"]]
        INC_MATRIZ.write_text(json.dumps(verdes, ensure_ascii=False, indent=1), encoding="utf-8")
        ban = json.loads(BANDERAS.read_text(encoding="utf-8"))
        ban["registros"] = [x for x in ban["registros"] if not str(x["id"]).startswith("MAT-")] + rojas
        BANDERAS.write_text(json.dumps(ban, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"escrito: {len(verdes)} incidencias verdes -> {INC_MATRIZ.name} | {len(rojas)} banderas rojas -> {BANDERAS.name}")


if __name__ == "__main__":
    main()
