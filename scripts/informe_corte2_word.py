"""Informes del corte 2 (mediodia) en Word.

Uso: python scripts/informe_corte2_word.py [hoja_respuestas.xlsx] [--descargar] [--resumen | --detalle]
  --resumen  informe breve para quienes toman decisiones (3 paginas)
  --detalle  informe detallado: departamentos, cada local, incidencias por tema y pendientes
  Sin opcion se generan los dos. --descargar baja la hoja desde data/entrada/forms_url.txt.
Recalcula las cifras con informe_corte2.datos() y escribe seguimiento/INFORME_*_CORTE2_Q20261004_HHMM.docx (carpeta ignorada por Git: lleva nombres).
Texto corrido y vinetas justificados; titulos, tablas y pies sin justificar.
"""
from __future__ import annotations

import random
import re
import sys
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import informe_corte2 as ic  # noqa: E402

AZUL = RGBColor(0x0B, 0x3D, 0x6E)
GRIS = RGBColor(0x5B, 0x67, 0x73)
DEP = {"Ancash": "Áncash", "Junin": "Junín", "Huanuco": "Huánuco", "San Martin": "San Martín", "Apurimac": "Apurímac"}
dep = lambda s: DEP.get(s, s)  # noqa: E731


def shade(cell, color: str) -> None:
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(sh)


def campo(par, instruccion: str) -> None:
    for tipo, texto in (("begin", None), (None, instruccion), ("end", None)):
        r = par.add_run()
        if tipo:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tipo)
            r._r.append(f)
        else:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = texto
            r._r.append(t)
        r.font.size = Pt(8.5)


class Informe:
    """Documento A4 con los estilos y ayudas comunes a los dos informes."""

    def __init__(self, d: dict, titulo: str, nota: str = "Documento de uso interno: contiene nombres de orientadores y monitores."):
        self.d = d
        self.doc = doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.1)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        st = doc.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        for n, sz in (("Heading 1", 14), ("Heading 2", 12)):
            s = doc.styles[n]
            s.font.name, s.font.size, s.font.bold, s.font.color.rgb = "Calibri", Pt(sz), True, AZUL
            s.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        self.p(titulo, size=17, bold=True, align=AL.CENTER, after=2, color=AZUL)
        self.p("Elecciones Regionales y Municipales 2026 · Domingo 4 de octubre de 2026", align=AL.CENTER, after=2)
        self.p(f"Corte 2 (mediodía) · datos al {d['corte_datos']}", bold=True, align=AL.CENTER, after=2, color=AZUL)
        self.p("Subdirección de Procedimiento Electoral y Georreferenciación (SDPEG) · Equipo de Datos", size=9.5, align=AL.CENTER, after=2, color=GRIS)
        self.p(nota, size=9, italic=True, align=AL.CENTER, after=8, color=GRIS)
        self.sec = sec

    def p(self, texto, size=None, align=AL.JUSTIFY, after=6, bold=False, italic=False, color=None):
        par = self.doc.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_after = Pt(after)
        r = par.add_run(texto)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        return par

    def h1(self, t):
        self.doc.add_heading(t, 1)

    def h2(self, t):
        self.doc.add_heading(t, 2)

    def tabla(self, cabecera, filas, anchos, size=9.5, centradas=(), colores=None):
        t = self.doc.add_table(rows=1, cols=len(cabecera))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(cabecera):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(size), RGBColor(255, 255, 255)
            shade(c, "0B3D6E")
        for n, fila in enumerate(filas):
            cells = t.add_row().cells
            for i, v in enumerate(fila):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                par.alignment = AL.CENTER if i in centradas else AL.LEFT
                r = par.add_run(str(v))
                r.font.size = Pt(size)
                if i == 0:
                    r.bold = True
                if colores and colores.get((n, i)):
                    shade(cells[i], colores[(n, i)])
        for fila in t.rows:
            for i, w in enumerate(anchos):
                fila.cells[i].width = Cm(w)
            fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        t.rows[0]._tr.get_or_add_trPr().append(th)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kpis(self, items):
        k = self.doc.add_table(rows=1, cols=len(items))
        k.style = "Table Grid"
        k.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (v, l, pc) in enumerate(items):
            c = k.rows[0].cells[i]
            c.width = Cm(16.8 / len(items))
            shade(c, "EAF0F7")
            c.text = ""
            a = c.paragraphs[0]
            a.alignment = AL.CENTER
            r = a.add_run(v)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(20), AZUL
            b = c.add_paragraph()
            b.alignment = AL.CENTER
            b.add_run(l).font.size = Pt(9)
            e = c.add_paragraph()
            e.alignment = AL.CENTER
            er = e.add_run(pc)
            er.bold, er.font.size = True, Pt(10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def incidencias(self, items, vacio="Sin incidencias reportadas en este corte."):
        if not items:
            self.p(vacio, italic=True, color=GRIS)
            return
        self.tabla(["Hora", "Distrito", "Local de votación", "Tipo", "Qué se reportó"],
                   [(x["hora"], f"{x['distrito']} ({dep(x['departamento'])})", x["local"] or "—", "Grave" if x["grave"] else "De atención", x["texto"]) for x in items],
                   [1.3, 3.2, 3.5, 1.9, 6.9], centradas=(0,))

    def guardar(self, salida: Path, etiqueta: str):
        self.p(f"Fuente: formulario de orientadores y matriz de seguimiento de la SDPEG. Datos al {self.d['corte_datos']}.", size=9, italic=True, color=GRIS, after=0)
        pie = self.sec.footer.paragraphs[0]
        pie.alignment = AL.CENTER
        pie.add_run(f"SDPEG · {etiqueta}, corte 2 · datos al {self.d['corte_datos']} · Página ").font.size = Pt(8.5)
        campo(pie, "PAGE")
        self.doc.save(salida)


def kpis_corte2(d: dict):
    pl, c1, c2 = d["plan"], d["c1"], d["c2"]
    return [(f"{c2['orientadores']} de {pl['orientadores']}", "orientadores con registro del corte 2", ic.pct(c2["orientadores"], pl["orientadores"])),
            (f"{c2['locales']} de {pl['locales']}", "locales con registro del corte 2", ic.pct(c2["locales"], pl["locales"])),
            (f"{c2['distritos']} de {pl['distritos']}", "distritos con registro del corte 2", ic.pct(c2["distritos"], pl["distritos"])),
            (f"{c1['orientadores']} de {pl['orientadores']}", "orientadores que asistieron (corte 1)", ic.pct(c1["orientadores"], pl["orientadores"]))]


def tabla_monitores(inf: Informe, d: dict):
    inf.tabla(["Monitor operativo", "Orientadores", "Ya registraron", "Faltan", "Avance"],
              [(m["monitor"], m["planificados"], len(m["c2_reg"]), len(m["c2_falta"]), ic.pct(len(m["c2_reg"]), m["planificados"])) for m in d["monitores"]],
              [5.0, 3.0, 3.0, 2.6, 3.0], centradas=(1, 2, 3, 4))
    inf.h2("Quiénes ya registraron y quiénes faltan")
    inf.tabla(["Monitor operativo", "Ya registraron", "Faltan"],
              [(m["monitor"], "\n".join(m["c2_reg"]) or "—", "\n".join(m["c2_falta"]) or "—") for m in d["monitores"]], [3.2, 6.8, 6.8], size=9)


def resumen(d: dict, salida: Path) -> None:
    inf = Informe(d, "INFORME DE AVANCE DE LA ORIENTACIÓN ELECTORAL")
    t = ic.parrafos(d)
    inf.h1("1. Resumen")
    inf.p(t["resumen"])
    inf.kpis(kpis_corte2(d))
    inf.h1("2. Avance del corte del mediodía por monitor operativo")
    inf.p(t["monitores"])
    tabla_monitores(inf, d)
    inf.h1("3. Incidencias reportadas")
    inf.p(t["incidencias"])
    inf.h2("Corte 1 · Llegada a los locales (08:00 aprox.)")
    inf.incidencias(d["incidencias"]["c1"])
    inf.h2("Corte 2 · Mediodía (12:00 aprox.)")
    inf.incidencias(d["incidencias"]["c2"])
    inf.h2("Otros reportes de la jornada (otros canales)")
    inf.incidencias(d["otros"])
    inf.guardar(salida, "Informe de avance")


def temas(d: dict) -> list[tuple[str, int]]:
    todas = d["incidencias"]["c1"] + d["incidencias"]["c2"] + d["otros"]
    reglas = [("Restitución o cambio de domicilio", r"restitu|domicilio|direccion"), ("Presuntos electores «golondrinos»", r"golondrin"),
              ("Apertura de mesas, miembros de mesa o documentos", r"mesa|dni"), ("Reclamo o malestar de ciudadanos", r"molest|exige|reclam|alter|protest|aglomer")]
    out = []
    for nombre, patron in reglas:
        n = sum(1 for x in todas if re.search(patron, ic.nz(x["texto"]).lower()))
        out.append((nombre, n))
    return out


def detalle(d: dict, salida: Path) -> None:
    inf = Informe(d, "INFORME DETALLADO DE AVANCE DE LA ORIENTACIÓN ELECTORAL")
    t = ic.parrafos(d)
    pl, c1, c2, locs = d["plan"], d["c1"], d["c2"], d["locales"]
    completos = [l for l in locs if l["c2"] >= l["planificados"]]
    parciales = [l for l in locs if 0 < l["c2"] < l["planificados"]]
    sin = [l for l in locs if l["c2"] == 0]
    todas = d["incidencias"]["c1"] + d["incidencias"]["c2"] + d["otros"]
    graves = sorted([x for x in todas if x["grave"]], key=lambda x: x["hora"])

    inf.h1("1. Resumen general")
    inf.p(t["resumen"])
    inf.kpis(kpis_corte2(d))
    inf.p(f"De los {pl['locales']} locales con orientadores, {len(completos)} tienen a todos sus orientadores registrados en el corte del mediodía, {len(parciales)} tienen un registro parcial "
          f"y {len(sin)} aún no registran. Estos datos corresponden al último envío recibido a las {d['hora']}; el corte seguirá recibiendo registros hasta su cierre.")

    inf.h1("2. Cobertura por departamento")
    deps = d["departamentos"]
    todos = [x["departamento"] for x in deps if x["orientadores"] and x["c2"] >= x["orientadores"]]
    rezag = sorted([x for x in deps if x["orientadores"] and x["c2"] < x["orientadores"]], key=lambda x: x["c2"] / x["orientadores"])
    frase = (f"{len(todos)} de los {len(deps)} departamentos tienen a todos sus orientadores registrados en el corte del mediodía. "
             + ("Los que aún tienen registros pendientes son " + ", ".join(f"{dep(x['departamento'])} ({x['c2']} de {x['orientadores']})" for x in rezag) + "." if rezag else "No hay departamentos con registros pendientes."))
    inf.p(frase)
    inf.tabla(["Departamento", "Locales", "Orientadores", "Corte 1", "Corte 2", "Avance corte 2", "Incidencias"],
              [(dep(x["departamento"]), x["locales"], x["orientadores"], x["c1"], x["c2"], ic.pct(x["c2"], x["orientadores"]), x["incidencias"]) for x in deps],
              [3.6, 2.0, 2.4, 2.0, 2.0, 2.8, 2.0], centradas=(1, 2, 3, 4, 5, 6),
              colores={(i, 5): ("D9EFE0" if x["c2"] >= x["orientadores"] else "FCEFC7") for i, x in enumerate(deps) if x["orientadores"]})

    inf.h1("3. Avance por monitor operativo")
    inf.p(t["monitores"])
    tabla_monitores(inf, d)

    inf.h1("4. Detalle por local de votación")
    inf.p("La tabla presenta los locales con orientadores, ordenados por departamento y distrito. La columna del corte 2 se resalta en verde cuando todos los orientadores del local "
          "ya registraron, en amarillo cuando el registro es parcial y en rojo cuando aún no hay ninguno.")
    colores = {}
    for i, l in enumerate(locs):
        colores[(i, 5)] = "D9EFE0" if l["c2"] >= l["planificados"] else "FCEFC7" if l["c2"] else "F8D7D3"
    inf.tabla(["Distrito (departamento)", "Local de votación", "Monitor operativo", "Orient.", "Corte 1", "Corte 2", "Incid.", "Falta registrar (corte 2)"],
              [(f"{l['distrito']} ({dep(l['departamento'])})", l["local"], ", ".join(l["monitores"]), l["planificados"], f"{l['c1']}/{l['planificados']}", f"{l['c2']}/{l['planificados']}",
                l["incidencias"] or "—", "\n".join(l["falta_c2"]) or "—") for l in locs],
              [2.9, 3.2, 2.3, 1.1, 1.3, 1.3, 1.1, 3.6], size=8, centradas=(3, 4, 5, 6), colores=colores)

    inf.h1("5. Incidencias reportadas")
    nc1, nc2, no = len(d["incidencias"]["c1"]), len(d["incidencias"]["c2"]), len(d["otros"])
    inf.p(f"Hasta el momento se han reportado {len(todas)} incidencias en la jornada: {nc1} en el corte de llegada, {nc2} en el corte del mediodía y {no} informadas por otros canales. "
          f"{len(graves)} son de mayor gravedad: están etiquetadas como graves o hacen referencia a violencia, y requieren seguimiento prioritario. Las demás son de atención y se resuelven en el propio local.")
    inf.h2("Incidencias que más resaltan")
    inf.incidencias(graves, "No hay incidencias graves reportadas.")
    inf.h2("Temas más frecuentes")
    tm = temas(d)
    mayor = max(tm, key=lambda x: x[1])
    inf.p(f"El tema que más se repite es «{mayor[0].lower()}», presente en {mayor[1]} de las {len(todas)} incidencias. Una misma incidencia puede corresponder a más de un tema.")
    inf.tabla(["Tema", "Incidencias"], tm, [12.0, 4.0], centradas=(1,))
    inf.h2("Corte 1 · Llegada a los locales (08:00 aprox.)")
    inf.incidencias(d["incidencias"]["c1"])
    inf.h2("Corte 2 · Mediodía (12:00 aprox.)")
    inf.incidencias(d["incidencias"]["c2"])
    inf.h2("Otros reportes de la jornada (otros canales)")
    inf.incidencias(d["otros"])

    inf.h1("6. Pendientes de registro")
    inf.p(f"Faltan por registrar {len(d['faltan_c2'])} orientadores en el corte del mediodía y {len(d['faltan_c1'])} en el corte de llegada. Se recomienda que cada monitor operativo "
          "confirme con ellos si ya se encuentran en su local y les solicite el registro.")
    filas = [("Corte 1", x["nombre"], x["local"], f"{x['distrito']} ({dep(x['departamento'])})", x["monitor"]) for x in d["faltan_c1"]]
    filas += [("Corte 2", x["nombre"], x["local"], f"{x['distrito']} ({dep(x['departamento'])})", x["monitor"]) for x in d["faltan_c2"]]
    inf.tabla(["Corte", "Orientador", "Local de votación", "Distrito (departamento)", "Monitor operativo"], filas, [1.6, 4.2, 4.0, 3.8, 3.2], size=9)
    inf.p("El corte 3, de cierre de la orientación, está previsto hacia las 16:30.")
    inf.guardar(salida, "Informe detallado")


def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    ruta = Path(args[0]) if args else ic.c.ENTRADA
    if "--descargar" in sys.argv:
        url = (ic.RAIZ / "data" / "entrada" / "forms_url.txt").read_text(encoding="utf-8").strip()
        urllib.request.urlretrieve(f"{url}&cb={random.randint(0, 10**9)}", ruta)
    d = ic.datos(ruta)
    ic.SEG.mkdir(exist_ok=True)
    hh = d["hora"].replace(":", "")
    for nombre, fn, quiero in (("RESUMEN", resumen, "--detalle" not in sys.argv), ("DETALLADO", detalle, "--resumen" not in sys.argv)):
        if quiero:
            salida = ic.SEG / (f"INFORME_AVANCE_CORTE2_Q20261004_{hh}.docx" if nombre == "RESUMEN" else f"INFORME_DETALLADO_CORTE2_Q20261004_{hh}.docx")
            fn(d, salida)
            print(f"{salida.name} | datos al {d['corte_datos']} | corte 2: {d['c2']['orientadores']} orientadores, {d['c2']['locales']} locales, {d['c2']['distritos']} distritos")


if __name__ == "__main__":
    main()
