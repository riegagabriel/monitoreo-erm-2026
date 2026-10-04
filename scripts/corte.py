"""Corte de la jornada: de la hoja de respuestas del formulario de orientadores a data/jornada.json (solo avance).

Uso: python scripts/corte.py [hoja_respuestas.xlsx]      (defecto: data/entrada/hoja_respuestas.xlsx)
Lee la pestaña RESP_ORIENTADORES por TITULO de columna (une los bloques repetidos con sufijo « 2»), cruza cada orientador con la base de
asignacion y cuenta, por local, cuantos orientadores reportaron llegada (corte 1) y cierre (corte 3). Las filas anteriores al inicio de la
jornada son ensayos y se excluyen. Un nombre que no este en la base aborta: nunca se descarta en silencio.
Salida: data/jornada.json, sin nombres ni monitores. Las incidencias NO se publican todavia (lista vacia).
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from actualizar_locales import BASE, HOJA, leer_adicionales, nz  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA = RAIZ / "data" / "entrada" / "hoja_respuestas.xlsx"
LOCALES = RAIZ / "data" / "orientacion_locales.json"
SALIDA = RAIZ / "data" / "jornada.json"
INICIO_JORNADA = dt.datetime(2026, 10, 4)
PESTANA = "RESP_ORIENTADORES"
T_NOMBRE = "Seleccione su nombre y local de votación."
T_CORTE = "¿Qué corte va a registrar?"
T_LLEGADA = "Hora de llegada al local de votación"
T_TERMINO = "Hora de término de la orientación electoral"


def _base_titulo(t) -> str:
    return re.sub(r"\s+\d+$", "", str(t or "").strip())


def _hora(v, fila: int, campo: str):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, dt.time):
        return v
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::\d{2})?", str(v).strip())
    if not m or int(m.group(1)) > 23:
        raise SystemExit(f"Fila {fila}: {campo} no es una hora HH:mm ({v!r})")
    return dt.time(int(m.group(1)), int(m.group(2)))


def normalizar(cabecera: list, filas: list) -> list[dict]:
    """Una fila de la hoja -> {fila, ts, k (nombre normalizado), corte (1-3), llegada, termino}. Une las columnas de titulo repetido."""
    cols = defaultdict(list)
    for i, t in enumerate(cabecera):
        cols[_base_titulo(t)].append(i)
    faltan = [t for t in (T_NOMBRE, T_CORTE, T_LLEGADA, T_TERMINO) if t not in cols]
    if faltan:
        raise SystemExit(f"No encuentro estos titulos en {PESTANA}: {faltan}")

    def val(r, t):
        return next((r[i] for i in cols[t] if r[i] not in (None, "")), None)

    out = []
    for n, r in enumerate(filas, start=2):
        if not any(c not in (None, "") for c in r):
            continue
        etiqueta, corte_t = val(r, T_NOMBRE), val(r, T_CORTE)
        m = re.search(r"Corte (\d)", str(corte_t or ""))
        if not etiqueta or not m:
            raise SystemExit(f"Fila {n}: falta el nombre o el corte")
        out.append({"fila": n, "ts": r[0], "k": nz(" ".join(str(etiqueta).split(" · ")[2:])), "corte": int(m.group(1)),
                    "llegada": _hora(val(r, T_LLEGADA), n, "llegada"), "termino": _hora(val(r, T_TERMINO), n, "termino")})
    return out


def avance_por_local(reg: list[dict], base: dict[str, list[str]], n_local: dict[str, int], inicio: dt.datetime):
    """base: nombre normalizado -> ids de sus locales (uno o mas); n_local: id -> orientadores en la base. Devuelve (avance, info)."""
    reales = [r for r in reg if r["ts"] >= inicio]
    desconocidos = sorted({r["k"] for r in reales if r["k"] not in base})
    if desconocidos:
        raise SystemExit(f"Orientadores que no estan en la base de asignacion: {desconocidos}")
    lleg, cier = defaultdict(set), defaultdict(set)
    for r in reales:
        for lid in base[r["k"]]:
            if r["corte"] == 1 and r["llegada"]:
                lleg[lid].add(r["k"])
            if r["corte"] == 3 and r["termino"]:
                cier[lid].add(r["k"])
    av = {}
    for lid, n in n_local.items():
        a, c = len(lleg[lid]), len(cier[lid])
        if a > n or c > n:
            raise SystemExit(f"Local {lid}: {a} llegadas y {c} cierres para {n} orientadores en la base")
        av[lid] = {"orientadores_base": n, "llegaron": a, "cierre": c}
    info = {"envios_reales": len(reales), "pruebas_excluidas": len(reg) - len(reales), "orientadores_con_llegada": len(set().union(*lleg.values())),
            "orientadores_con_cierre": len(set().union(*cier.values())), "ultimo_envio": max(r["ts"] for r in reales) if reales else None}
    return av, info


def cargar_base(locales: dict) -> dict[str, list[str]]:
    """nombre normalizado -> ids de local, desde la base de asignacion mas las asignaciones adicionales (llave del local: distrito + nombre)."""
    ws = openpyxl.load_workbook(BASE, data_only=True)[HOJA]
    cab = [str(c).strip() if c else "" for c in next(ws.iter_rows(values_only=True))]
    ix = {h: i for i, h in enumerate(cab)}
    por_local = {(nz(r["distrito"]), nz(r["local"])): r["id"] for r in locales["registros"]}
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        try:
            int(str(r[ix["N°"]]).strip())
        except ValueError:
            continue
        nombre = r[ix["NOMBRES Y APELLIDOS"]]
        if not nombre:
            continue
        k = nz(nombre)
        lid = por_local.get((nz(r[ix["DISTRITO"]]), nz(r[ix["NOMBRE DEL LOCAL"]])))
        if lid is None:
            raise SystemExit(f"Local de la base sin id en orientacion_locales.json: {r[ix['DISTRITO']]} · {r[ix['NOMBRE DEL LOCAL']]}")
        if k in out:
            raise SystemExit(f"Nombre repetido en la base: {k}")
        out[k] = [lid]
    for a in leer_adicionales():
        k, lid = nz(a["nombre"]), por_local.get((nz(a["distrito"]), nz(a["local"])))
        if k not in out or lid is None:
            raise SystemExit(f"Asignacion adicional sin orientador en la base o sin local en orientacion_locales.json: {a}")
        if lid not in out[k]:
            out[k].append(lid)
    return out


def main() -> None:
    ruta = Path(sys.argv[1]) if len(sys.argv) > 1 else ENTRADA
    locales = json.loads(LOCALES.read_text(encoding="utf-8"))
    base = cargar_base(locales)
    n_local = {r["id"]: r["n_orientadores"] for r in locales["registros"]}
    puestos = sum(len(v) for v in base.values())
    if sum(n_local.values()) != puestos:
        raise SystemExit(f"orientacion_locales.json suma {sum(n_local.values())} puestos y la base {puestos}: ejecute actualizar_locales.py")
    filas = list(openpyxl.load_workbook(ruta, data_only=True)[PESTANA].iter_rows(values_only=True))
    reg = normalizar(list(filas[0]), filas[1:])
    av, info = avance_por_local(reg, base, n_local, INICIO_JORNADA)
    ultimo = info["ultimo_envio"]
    completos = sum(1 for a in av.values() if a["llegaron"] == a["orientadores_base"])
    jornada = {"corte": ultimo.strftime("%d/%m/%Y %H:%M") if ultimo else "", "es_ejemplo": False,
               "fuente": "Formulario de orientadores (corte 1: llegada; corte 3: cierre)", "orientadores_base": len(base),
               "orientadores_con_llegada": info["orientadores_con_llegada"], "orientadores_con_cierre": info["orientadores_con_cierre"],
               "locales_completos": completos, "avance": av, "incidencias": []}
    SALIDA.write_text(json.dumps(jornada, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"corte {jornada['corte']} | envios reales {info['envios_reales']} (pruebas excluidas: {info['pruebas_excluidas']})")
    print(f"llegada reportada: {info['orientadores_con_llegada']} de {len(base)} orientadores | cierre: {info['orientadores_con_cierre']}")
    print(f"locales con todos sus orientadores: {completos} de {len(av)} -> {SALIDA}")


if __name__ == "__main__":
    main()
