"""Cifras del informe del corte 2 (mediodia) y su boceto en HTML.

Uso: python scripts/informe_corte2.py [hoja_respuestas.xlsx]     (defecto: data/entrada/hoja_respuestas.xlsx)
Lee la hoja de respuestas del formulario de orientadores, la base de asignacion y las incidencias ya validadas (incidencias_forms.csv e incidencias_matriz.json).
Escribe seguimiento/informe_corte2_datos.json y seguimiento/boceto_informe_corte2.html (carpeta ignorada por Git: lleva nombres de orientadores).
"""
from __future__ import annotations

import csv
import html
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import corte as c  # noqa: E402
from actualizar_locales import BASE, HOJA, nz  # noqa: E402

RAIZ = c.RAIZ
SEG = RAIZ / "seguimiento"
MENORES = {"de", "del", "la", "las", "los", "y", "e"}
VIOL = re.compile(r"\b(violen|amenaz|agresi|agred|golpe|pelea|rina\b|enfrentamiento|disturbio|atac|empujon|incendi|quem[oa]|destrozo|balacera|disparo|herid|lesion|intimid|hostig)")


def es_grave(texto: str, tipo_a: bool) -> bool:
    """Grave: etiquetada como grave (tipo A) o con referencia a violencia en el texto (misma regla del dashboard)."""
    return bool(tipo_a) or bool(VIOL.search(nz(texto).lower()))


def titulo(s: str) -> str:
    return " ".join(w.lower() if i and w.lower() in MENORES else w.capitalize() for i, w in enumerate(str(s).split()))


def leer_base() -> list[dict]:
    ws = openpyxl.load_workbook(BASE, data_only=True)[HOJA]
    cab = [str(x).strip() if x else "" for x in next(ws.iter_rows(values_only=True))]
    ix = {h: i for i, h in enumerate(cab)}
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        try:
            int(str(r[ix["N°"]]).strip())
        except ValueError:
            continue
        if r[ix["NOMBRES Y APELLIDOS"]]:
            out.append({"nombre": " ".join(str(r[ix["NOMBRES Y APELLIDOS"]]).split()), "monitor": " ".join(str(r[ix["MONITOR OPERATIVO"]]).split()),
                        "distrito": str(r[ix["DISTRITO"]]).strip(), "local": str(r[ix["NOMBRE DEL LOCAL"]]).strip()})
    return out


def datos(ruta: Path) -> dict:
    locales = json.loads(c.LOCALES.read_text(encoding="utf-8"))
    por_id = {r["id"]: r for r in locales["registros"]}
    por_local = {(nz(r["distrito"]), nz(r["local"])): r["id"] for r in locales["registros"]}
    base_lids = c.cargar_base(locales)
    filas = list(openpyxl.load_workbook(ruta, data_only=True)[c.PESTANA].iter_rows(values_only=True))
    reg = [r for r in c.normalizar(list(filas[0]), filas[1:]) if r["ts"] >= c.INICIO_JORNADA]
    c1 = {r["k"] for r in reg if r["corte"] == 1 and r["llegada"]}
    c2 = {r["k"] for r in reg if r["corte"] == 2}
    base = leer_base()
    mon = defaultdict(lambda: {"c1_reg": [], "c1_falta": [], "c2_reg": [], "c2_falta": []})
    for b in base:
        k = nz(b["nombre"])
        m = mon[b["monitor"]]
        m["c1_reg" if k in c1 else "c1_falta"].append(titulo(b["nombre"]))
        m["c2_reg" if k in c2 else "c2_falta"].append(titulo(b["nombre"]))
    monitores = [{"monitor": titulo(n), "planificados": len(m["c2_reg"]) + len(m["c2_falta"]), **{k: sorted(v) for k, v in m.items()}} for n, m in sorted(mon.items())]
    reg_c1 = defaultdict(set)
    reg_c2 = defaultdict(set)
    for r in reg:
        if (r["corte"] == 1 and r["llegada"]) or r["corte"] == 2:
            for lid in c._lids(r, base_lids, por_local):
                (reg_c1 if r["corte"] == 1 else reg_c2)[lid].add(r["k"])
    lids_c1, lids_c2 = set(reg_c1), set(reg_c2)
    persona_por_lid = defaultdict(set)
    for k, lids in base_lids.items():
        for lid in lids:
            persona_por_lid[lid].add(k)
    completos = sum(1 for lid, ks in persona_por_lid.items() if ks <= reg_c2[lid])
    dist = lambda ls: {por_id[l]["ubigeo_inei"] for l in ls}
    ult = max(r["ts"] for r in reg)

    inc = {1: [], 2: [], 3: []}
    if c.INC_FORMS.exists():
        with c.INC_FORMS.open(encoding="utf-8", newline="") as f:
            val = {x["id_envio"].strip(): x for x in csv.DictReader(f) if x["validado"].strip().upper() == "SI"}
        for r in reg:
            if r["inc"] != "Sí":
                continue
            lid = por_local[(r["dist"], r["local"])]
            x = val.get(f"ORI-{r['ts']:%Y%m%d%H%M%S}-{lid}")
            if x:
                l = por_id[lid]
                texto = x["resumen_publicable"].strip()
                inc[r["corte"]].append({"hora": x["hora"], "lid": lid, "u": l["ubigeo_inei"], "departamento": titulo(l["departamento"]), "distrito": titulo(l["distrito"]), "local": l["local"],
                                        "grave": es_grave(texto, x["tipo"].strip().upper() == "A"), "texto": texto})
    otros = []
    if c.INC_MATRIZ.exists():
        geo = {f["properties"]["u"]: f["properties"] for f in json.loads((RAIZ / "dashboard" / "public" / "data" / "precarga" / "distritos_pais.geojson").read_text(encoding="utf-8"))["features"]}
        for x in json.loads(c.INC_MATRIZ.read_text(encoding="utf-8")):
            g = geo[x["u"]]
            otros.append({"hora": f"{int(x['h']):02d}:{round((x['h'] % 1) * 60):02d}", "lid": x["l"], "u": x["u"], "departamento": titulo(g["dep"]), "distrito": titulo(g["dist"]),
                          "local": por_id[x["l"]]["local"] if x["l"] else "", "grave": es_grave(x["t"], x["g"] == 2), "texto": x["t"]})
    otros.sort(key=lambda x: x["hora"])
    for v in inc.values():
        v.sort(key=lambda x: x["hora"])
    nombres = {nz(b["nombre"]): titulo(b["nombre"]) for b in base}
    mon_persona = {nz(b["nombre"]): titulo(b["monitor"]) for b in base}
    todas = [x for v in inc.values() for x in v] + otros
    n_inc = defaultdict(int)
    for x in todas:
        if x["lid"]:
            n_inc[x["lid"]] += 1
    locs = []
    for lid, ks in persona_por_lid.items():
        l = por_id[lid]
        falta = [nombres[k] + (" (registró en otro local)" if k in c2 else "") for k in sorted(ks - reg_c2[lid], key=lambda k: nombres[k])]
        locs.append({"id": lid, "departamento": titulo(l["departamento"]), "provincia": titulo(l["provincia"]), "distrito": titulo(l["distrito"]), "local": l["local"],
                     "monitores": sorted({mon_persona[k] for k in ks}), "planificados": len(ks), "c1": len(ks & reg_c1[lid]), "c2": len(ks & reg_c2[lid]),
                     "falta_c2": falta, "incidencias": n_inc[lid]})
    locs.sort(key=lambda x: (x["departamento"], x["distrito"], x["local"]))
    deps = {}
    for k, lids in base_lids.items():
        x = deps.setdefault(titulo(por_id[lids[0]]["departamento"]), {"orientadores": 0, "c1": 0, "c2": 0, "locales": set(), "incidencias": 0})
        x["orientadores"] += 1
        x["c1"] += k in c1
        x["c2"] += k in c2
    for l in locs:
        deps[l["departamento"]]["locales"].add(l["id"])
    for x in todas:
        deps.setdefault(x["departamento"], {"orientadores": 0, "c1": 0, "c2": 0, "locales": set(), "incidencias": 0})["incidencias"] += 1
    departamentos = sorted(({"departamento": n, **{k: (len(v) if k == "locales" else v) for k, v in x.items()}} for n, x in deps.items()), key=lambda x: x["departamento"])

    def faltan(hechos):
        out = []
        for k, lids in base_lids.items():
            if k not in hechos:
                l = por_id[lids[0]]
                out.append({"nombre": nombres[k], "local": l["local"], "distrito": titulo(l["distrito"]), "departamento": titulo(l["departamento"]), "monitor": mon_persona[k]})
        return sorted(out, key=lambda x: (x["monitor"], x["nombre"]))

    return {"corte_datos": ult.strftime("%d/%m/%Y %H:%M"), "hora": ult.strftime("%H:%M"),
            "locales": locs, "departamentos": departamentos, "faltan_c1": faltan(c1), "faltan_c2": faltan(c2),
            "plan": {"orientadores": len(base), "locales_base": 28, "locales": len(por_id), "distritos": len({r["ubigeo_inei"] for r in por_id.values()})},
            "c1": {"orientadores": len(c1), "locales": len(lids_c1)},
            "c2": {"orientadores": len(c2), "locales": len(lids_c2), "distritos": len(dist(lids_c2)), "locales_completos": completos},
            "monitores": monitores, "incidencias": {"c1": inc[1], "c2": inc[2]}, "otros": otros}


def pct(a: int, b: int) -> str:
    return f"{round(100 * a / b)} %" if b else "0 %"


def parrafos(d: dict) -> dict[str, str]:
    p, c1, c2 = d["plan"], d["c1"], d["c2"]
    ms = sorted(d["monitores"], key=lambda m: len(m["c2_reg"]) / m["planificados"], reverse=True)
    mejor, peor = ms[0], ms[-1]
    ni1, ni2 = len(d["incidencias"]["c1"]), len(d["incidencias"]["c2"])
    graves = [x for x in d["incidencias"]["c1"] + d["incidencias"]["c2"] + d["otros"] if x["grave"]]
    res = (f"De un total de {p['orientadores']} orientadores planificados, {c1['orientadores']} asistieron al inicio de la jornada. Se planificaron {p['locales_base']} locales "
           f"con orientadores y se añadió 1 en Pichari (IE Parque Industrial), con lo que la orientación se brinda en {p['locales']} locales de {p['distritos']} distritos. "
           f"A las {d['hora']}, en el corte del mediodía, han registrado su avance {c2['orientadores']} orientadores ({pct(c2['orientadores'], p['orientadores'])}), "
           f"que atienden {c2['locales']} locales de {p['locales']} ({pct(c2['locales'], p['locales'])}) en {c2['distritos']} distritos de {p['distritos']}.")
    mon = (f"Por monitor operativo, el mayor avance corresponde a {mejor['monitor']}, con {len(mejor['c2_reg'])} de {mejor['planificados']} orientadores registrados, "
           f"y el menor a {peor['monitor']}, con {len(peor['c2_reg'])} de {peor['planificados']}. Faltan por registrar el corte del mediodía "
           f"{p['orientadores'] - c2['orientadores']} orientadores; el detalle por monitor se presenta en la segunda tabla.")
    inc = (f"Hasta el momento se han reportado {ni1 + ni2 + len(d['otros'])} incidencias en la jornada: {ni1} en el corte de llegada, {ni2} en el corte del mediodía y "
           f"{len(d['otros'])} informadas por otros canales. {len(graves)} de ellas son de mayor gravedad y requieren seguimiento prioritario.")
    return {"resumen": res, "monitores": mon, "incidencias": inc}


def tabla_inc(items: list[dict]) -> str:
    if not items:
        return "<p class='vacio'>Sin incidencias reportadas en este corte.</p>"
    f = "".join(f"<tr><td>{html.escape(x['hora'])}</td><td><b>{html.escape(x['distrito'])}</b><br><span class='mut'>{html.escape(x['departamento'])}</span></td>"
                f"<td>{html.escape(x['local'])}</td><td><span class='tag {'g' if x['grave'] else 'l'}'>{'Grave' if x['grave'] else 'De atención'}</span></td>"
                f"<td>{html.escape(x['texto'])}</td></tr>" for x in items)
    return f"<table><thead><tr><th>Hora</th><th>Distrito</th><th>Local de votación</th><th>Tipo</th><th>Qué se reportó</th></tr></thead><tbody>{f}</tbody></table>"


def boceto(d: dict) -> str:
    p, c1, c2, t = d["plan"], d["c1"], d["c2"], parrafos(d)
    kpi = "".join(f"<div class='kpi'><b>{v}</b><span>{html.escape(l)}</span><small>{html.escape(s)}</small></div>" for v, l, s in [
        (f"{c2['orientadores']} de {p['orientadores']}", "orientadores con registro del corte 2", pct(c2['orientadores'], p['orientadores'])),
        (f"{c2['locales']} de {p['locales']}", "locales con registro del corte 2", pct(c2['locales'], p['locales'])),
        (f"{c2['distritos']} de {p['distritos']}", "distritos con registro del corte 2", pct(c2['distritos'], p['distritos'])),
        (f"{c1['orientadores']} de {p['orientadores']}", "orientadores que asistieron (corte 1)", pct(c1['orientadores'], p['orientadores']))])
    res = "".join(f"<tr><td><b>{html.escape(m['monitor'])}</b></td><td>{m['planificados']}</td><td>{len(m['c2_reg'])}</td><td>{len(m['c2_falta'])}</td>"
                  f"<td><div class='bar'><i style='width:{100 * len(m['c2_reg']) / m['planificados']:.0f}%'></i></div></td><td>{pct(len(m['c2_reg']), m['planificados'])}</td></tr>" for m in d["monitores"])
    det = "".join(f"<tr><td><b>{html.escape(m['monitor'])}</b></td><td>{'<br>'.join(map(html.escape, m['c2_reg'])) or '—'}</td><td>{'<br>'.join(map(html.escape, m['c2_falta'])) or '—'}</td></tr>" for m in d["monitores"])
    return f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Boceto · Informe del corte 2</title>
<style>
:root{{--tinta:#1c2733;--mut:#5b6773;--linea:#d6dbe0;--marca:#0b3d6e;--fondo:#eef1f4;--ok:#1d7a4a;--ger:#b42318}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--fondo);font:15px/1.55 Calibri,'Segoe UI',Arial,sans-serif;color:var(--tinta)}}
.aviso{{background:#fff4d6;border-bottom:1px solid #e6c766;padding:10px 20px;font-size:13px}}
.hoja{{max-width:860px;margin:24px auto;background:#fff;padding:48px 56px;box-shadow:0 1px 6px #0002}}
h1{{font-size:24px;margin:0 0 4px;color:var(--marca)}}.sub{{color:var(--mut);margin:0 0 4px}}.corte{{display:inline-block;margin:6px 0 22px;padding:3px 10px;background:#e8eef6;color:var(--marca);border-radius:4px;font-size:13px;font-weight:600}}
h2{{font-size:17px;margin:28px 0 8px;color:var(--marca);border-bottom:2px solid var(--marca);padding-bottom:3px}}
p{{text-align:justify;margin:8px 0}}.kpis{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:14px 0}}
.kpi{{border:1px solid var(--linea);border-top:3px solid var(--marca);padding:10px 12px}}.kpi b{{display:block;font-size:24px;color:var(--marca)}}.kpi span{{display:block;font-size:12.5px;line-height:1.3}}.kpi small{{color:var(--mut);font-size:12px}}
table{{width:100%;border-collapse:collapse;margin:10px 0;font-size:13.5px}}th{{background:var(--marca);color:#fff;text-align:left;padding:6px 8px;font-weight:600}}td{{padding:6px 8px;border-bottom:1px solid var(--linea);vertical-align:top}}
.mut{{color:var(--mut);font-size:12px}}.bar{{height:9px;background:#e3e8ee;border-radius:5px;min-width:80px}}.bar i{{display:block;height:100%;background:var(--ok);border-radius:5px}}
.tag{{padding:1px 8px;border-radius:10px;font-size:12px;font-weight:600}}.tag.g{{background:#fde4e1;color:var(--ger)}}.tag.l{{background:#e6eef8;color:var(--marca)}}
.vacio{{color:var(--mut);font-style:italic}}h3{{font-size:14.5px;margin:16px 0 4px}}.pie{{margin-top:30px;font-size:12px;color:var(--mut);border-top:1px solid var(--linea);padding-top:8px}}
@media(max-width:700px){{.hoja{{padding:24px 18px}}.kpis{{grid-template-columns:repeat(2,1fr)}}}}
</style></head><body>
<div class="aviso"><b>BOCETO</b> · cifras reales al {d['corte_datos']} con el corte 2 todavía en curso; el informe final se genera al cierre. Esta franja no forma parte del documento.</div>
<main class="hoja">
<h1>Informe de avance de la orientación electoral</h1>
<p class="sub">Elecciones Regionales y Municipales 2026 · Domingo 4 de octubre de 2026</p>
<span class="corte">Corte 2 (mediodía) · datos al {d['corte_datos']}</span>
<h2>1. Resumen</h2>
<p>{html.escape(t['resumen'])}</p>
<div class="kpis">{kpi}</div>
<h2>2. Avance del corte del mediodía por monitor operativo</h2>
<p>{html.escape(t['monitores'])}</p>
<table><thead><tr><th>Monitor operativo</th><th>Orientadores</th><th>Ya registraron</th><th>Faltan</th><th>Avance</th><th></th></tr></thead><tbody>{res}</tbody></table>
<h3>Quiénes ya registraron y quiénes faltan</h3>
<table><thead><tr><th style="width:17%">Monitor operativo</th><th>Ya registraron</th><th>Faltan</th></tr></thead><tbody>{det}</tbody></table>
<h2>3. Incidencias reportadas</h2>
<p>{html.escape(t['incidencias'])}</p>
<h3>Corte 1 · Llegada a los locales (08:00 aprox.)</h3>{tabla_inc(d['incidencias']['c1'])}
<h3>Corte 2 · Mediodía (12:00 aprox.)</h3>{tabla_inc(d['incidencias']['c2'])}
<h3>Otros reportes de la jornada (otros canales)</h3>{tabla_inc(d['otros'])}
<p class="pie">Fuente: formulario de orientadores y matriz de seguimiento de la SDPEG. Datos al {d['corte_datos']}.</p>
</main></body></html>"""


def main() -> None:
    ruta = Path(sys.argv[1]) if len(sys.argv) > 1 else c.ENTRADA
    d = datos(ruta)
    SEG.mkdir(exist_ok=True)
    (SEG / "informe_corte2_datos.json").write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
    (SEG / "boceto_informe_corte2.html").write_text(boceto(d), encoding="utf-8")
    print(f"datos al {d['corte_datos']} | corte 2: {d['c2']['orientadores']} orientadores, {d['c2']['locales']} locales, {d['c2']['distritos']} distritos | "
          f"incidencias c1 {len(d['incidencias']['c1'])}, c2 {len(d['incidencias']['c2'])}, otros {len(d['otros'])}")


if __name__ == "__main__":
    main()
