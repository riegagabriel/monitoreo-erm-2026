"""Informe ejecutivo para autoridades: cifras clave, orientacion brindada y situacion de los casos graves. Boceto en HTML y version en Word.

Uso: python scripts/informe_ejecutivo.py [hoja_respuestas.xlsx] [--descargar] [--boceto | --word]
  Sin opcion genera las dos salidas. Lee el formulario de orientadores, las incidencias validadas y data/entrada/contexto_casos.csv
  (id, distritos, texto): informacion de campo sobre los casos graves, que se escribe a mano y se imprime tal cual en el informe.
Escribe seguimiento/boceto_informe_ejecutivo.html y seguimiento/INFORME_EJECUTIVO_Q20261004_HHMM.docx (carpeta ignorada por Git).
No lleva nombres de orientadores.
"""
from __future__ import annotations

import csv
import html
import json
import random
import sys
import urllib.request
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import corte as c  # noqa: E402
import informe_corte2 as ic  # noqa: E402
import informe_corte2_word as iw  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import RGBColor  # noqa: E402

CONTEXTO = c.RAIZ / "data" / "entrada" / "contexto_casos.csv"
DEP = {"Ancash": "Áncash", "Junin": "Junín", "Huanuco": "Huánuco", "San Martin": "San Martín", "Apurimac": "Apurímac"}
dep = lambda s: DEP.get(s, s)  # noqa: E731
DIST = {"Masin": "Masín", "Chavin de Huantar": "Chavín de Huántar"}
dist = lambda s: DIST.get(s, s)  # noqa: E731


def datos(ruta: Path) -> dict:
    d = ic.datos(ruta)
    filas = list(openpyxl.load_workbook(ruta, data_only=True)[c.PESTANA].iter_rows(values_only=True))
    reg = c.normalizar(list(filas[0]), filas[1:])
    d["consultas"] = c.resumen_consultas(reg, c.INICIO_JORNADA)
    d["contexto"] = []
    if CONTEXTO.exists():
        with CONTEXTO.open(encoding="utf-8", newline="") as f:
            d["contexto"] = [{"distritos": x["distritos"].strip(), "texto": x["texto"].strip()} for x in csv.DictReader(f)]
    todas = d["incidencias"]["c1"] + d["incidencias"]["c2"] + d["otros"]
    for x in todas:
        x["distrito"] = dist(x["distrito"])
    d["todas"] = todas
    d["graves"] = sorted([x for x in todas if x["grave"]], key=lambda x: x["hora"], reverse=True)
    riesgo = {x["ubigeo_inei"]: (x.get("restituidos") or {}) for x in json.loads((c.RAIZ / "dashboard" / "public" / "data" / "precarga" / "riesgo_previo.json").read_text(encoding="utf-8"))["distritos"]}
    for x in todas:
        x["rest"] = riesgo.get(x["u"], {}).get("total", 0)
        x["rest_reniec"], x["rest_jne"] = riesgo.get(x["u"], {}).get("reniec", 0), riesgo.get(x["u"], {}).get("jne", 0)
    dist_inc = {x["u"] for x in todas}
    dist_gra = {x["u"] for x in d["graves"]}
    suma = lambda us, k: sum(riesgo.get(u, {}).get(k, 0) for u in us)  # noqa: E731
    d["rest_inc"] = {"distritos": len(dist_inc), "total": suma(dist_inc, "total"), "reniec": suma(dist_inc, "reniec"), "jne": suma(dist_inc, "jne"),
                     "graves_distritos": len(dist_gra), "graves_total": suma(dist_gra, "total"), "graves_reniec": suma(dist_gra, "reniec"), "graves_jne": suma(dist_gra, "jne"), "graves_sin": sum(1 for u in dist_gra if not riesgo.get(u, {}).get("total"))}
    d["deps_completos"] = sum(1 for x in d["departamentos"] if x["orientadores"] and x["c2"] >= x["orientadores"])
    return d


def pc(n: int, total: int) -> str:
    p = 100 * n / total if total else 0
    return "<1 %" if 0 < p < 1 else f"{round(p)} %"


def fn(n: int) -> str:
    return f"{n:,}".replace(",", " ")


def lista(items: list[str]) -> str:
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " y " + items[-1]


def textos(d: dict) -> dict[str, str]:
    p, c1, c2, cs = d["plan"], d["c1"], d["c2"], d["consultas"]
    ndep = len(d["departamentos"])
    general = (f"Al corte de las {d['hora']}, {c1['orientadores']} de los {p['orientadores']} orientadores planificados asistieron a sus locales ({pc(c1['orientadores'], p['orientadores'])}), "
               f"en {p['locales']} locales de votación de {p['distritos']} distritos y {ndep} departamentos. En el corte del mediodía han registrado su avance {c2['orientadores']} orientadores "
               f"({pc(c2['orientadores'], p['orientadores'])}) y {c2['locales_completos']} de los {p['locales']} locales tienen a todo su personal registrado.")
    top = [(t, n) for t, n in cs["por_tipo"].items() if n > 0][:3]
    nm = lambda t: t if t.startswith("DNI") else t[0].lower() + t[1:]  # noqa: E731
    orient = (f"En total consultaron {cs['total']:,} personas".replace(",", " ") + (". La mayor parte lo hizo por " + lista([f"{nm(t)} ({n}; {pc(n, cs['total'])})" for t, n in top]) + "." if top else ".")
              + " Las cifras son estimaciones de cada orientador para toda la jornada.")
    g = d["graves"]
    inc = (f"Se han reportado {len(d['todas'])} incidencias durante la jornada, de las cuales {len(g)} son graves: se registraron en "
           + lista([f"{x['distrito']} ({dep(x['departamento'])})" for x in sorted(g, key=lambda x: x["hora"])])
           + f". Las otras {len(d['todas']) - len(g)} son de atención.")
    r = d["rest_inc"]
    nombres_sin = sorted({x["distrito"] for x in d["graves"] if not x["rest"]})
    sin = f" En {len(nombres_sin)} de ellos ({lista(nombres_sin)}) no hay restituidos registrados." if nombres_sin else ""
    rest = (f"Las incidencias se registraron en {r['distritos']} distritos, donde hay {fn(r['total'])} ciudadanos restituidos ({fn(r['reniec'])} por RENIEC y {fn(r['jne'])} por el JNE). "
            f"En los {r['graves_distritos']} distritos con incidencias graves suman {fn(r['graves_total'])} ({fn(r['graves_reniec'])} por RENIEC y {fn(r['graves_jne'])} por el JNE)." + sin)
    return {"general": general, "orientacion": orient, "incidencias": inc, "restituidos": rest}


def kpis(d: dict) -> list[tuple[str, str, str]]:
    p, c1, c2, cs = d["plan"], d["c1"], d["c2"], d["consultas"]
    return [(f"{c1['orientadores']} de {p['orientadores']}", "orientadores asistieron a su local", pc(c1["orientadores"], p["orientadores"])),
            (f"{p['locales']}", "locales de votación con orientación", f"en {p['distritos']} distritos"),
            (f"{c2['orientadores']} de {p['orientadores']}", "orientadores con registro del mediodía", pc(c2["orientadores"], p["orientadores"])),
            (f"{cs['total']:,}".replace(",", " "), "personas consultaron (estimado)", f"{cs['orientadores']} orientadores reportaron"),
            (f"{len(d['todas'])}", "incidencias reportadas", f"{len(d['graves'])} graves · {len(d['todas']) - len(d['graves'])} de atención"),
            (f"{len({x['distrito'] for x in d['graves']})}", "distritos con incidencias graves", "con seguimiento en curso")]


def boceto(d: dict) -> str:
    t, cs = textos(d), d["consultas"]
    tarj = "".join(f"<div class='kpi'><b>{html.escape(v)}</b><span>{html.escape(l)}</span><small>{html.escape(s)}</small></div>" for v, l, s in kpis(d))
    barras = "".join(f"<div class='gf'><span>{html.escape(k)}</span><span class='pista'><i style='width:{100 * n / max(1, list(cs['por_tipo'].values())[0]):.0f}%'></i></span><b>{n}</b><em>{pc(n, cs['total'])}</em></div>"
                     for k, n in cs["por_tipo"].items() if n > 0)
    ctx = "".join(f"<div class='ctx'><b>{html.escape(x['distritos'])}</b><p>{html.escape(x['texto'])}</p></div>" for x in d["contexto"])
    filas = "".join(f"<tr><td>{html.escape(x['hora'])}</td><td><b>{html.escape(x['distrito'])}</b><br><span class='mut'>{html.escape(dep(x['departamento']))}</span></td><td class='num'>{fn(x['rest_reniec']) if x['rest'] else '—'}</td><td class='num'>{fn(x['rest_jne']) if x['rest'] else '—'}</td><td>{html.escape(x['texto'])}</td></tr>" for x in d["graves"])
    return f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Boceto · Informe ejecutivo</title>
<style>
:root{{--tinta:#1c2733;--mut:#5b6773;--linea:#d6dbe0;--marca:#0b3d6e;--fondo:#eef1f4;--ger:#b42318}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--fondo);font:15px/1.55 Calibri,'Segoe UI',Arial,sans-serif;color:var(--tinta)}}
.aviso{{background:#fff4d6;border-bottom:1px solid #e6c766;padding:10px 20px;font-size:13px}}
.hoja{{max-width:860px;margin:24px auto;background:#fff;padding:44px 56px;box-shadow:0 1px 6px #0002}}
h1{{font-size:23px;margin:0 0 4px;color:var(--marca)}}.sub{{color:var(--mut);margin:0 0 4px}}.corte{{display:inline-block;margin:6px 0 20px;padding:3px 10px;background:#e8eef6;color:var(--marca);border-radius:4px;font-size:13px;font-weight:600}}
h2{{font-size:16.5px;margin:24px 0 8px;color:var(--marca);border-bottom:2px solid var(--marca);padding-bottom:3px}}p{{text-align:justify;margin:8px 0}}
.kpis{{display:grid;grid-template-columns:repeat(3,1fr);gap:10px;margin:14px 0}}.kpi{{border:1px solid var(--linea);border-top:3px solid var(--marca);padding:10px 12px}}
.kpi b{{display:block;font-size:26px;color:var(--marca)}}.kpi span{{display:block;font-size:12.5px;line-height:1.3}}.kpi small{{color:var(--mut);font-size:12px}}
.gf{{display:grid;grid-template-columns:200px 1fr 44px 48px;gap:10px;align-items:center;margin:6px 0;font-size:13.5px}}.pista{{height:10px;background:#e3e8ee;border-radius:5px}}.pista i{{display:block;height:100%;background:var(--marca);border-radius:5px}}.gf em{{font-style:normal;color:var(--mut);text-align:right}}
.ctx{{border-left:4px solid var(--ger);background:#fdf2f1;padding:8px 14px;margin:10px 0}}.ctx b{{color:var(--ger)}}
table{{width:100%;border-collapse:collapse;margin:10px 0;font-size:13.5px}}th{{background:var(--marca);color:#fff;text-align:left;padding:6px 8px}}td{{padding:6px 8px;border-bottom:1px solid var(--linea);vertical-align:top}}.mut{{color:var(--mut);font-size:12px}}td.num{{text-align:right;white-space:nowrap}}
.pie{{margin-top:26px;font-size:12px;color:var(--mut);border-top:1px solid var(--linea);padding-top:8px}}
@media(max-width:700px){{.hoja{{padding:24px 18px}}.kpis{{grid-template-columns:repeat(2,1fr)}}.gf{{grid-template-columns:1fr 44px 48px}}.gf .pista{{grid-column:1/-1}}}}
</style></head><body>
<div class="aviso"><b>BOCETO</b> · cifras reales al {d['corte_datos']}, con el corte del mediodía en curso. Esta franja no forma parte del documento.</div>
<main class="hoja">
<h1>Informe ejecutivo de la orientación electoral</h1>
<p class="sub">Elecciones Regionales y Municipales 2026 · Domingo 4 de octubre de 2026</p>
<span class="corte">Datos al {d['corte_datos']}</span>
<h2>1. Situación general</h2><p>{html.escape(t['general'])}</p><div class="kpis">{tarj}</div>
<h2>2. Orientación brindada</h2><p>{html.escape(t['orientacion'])}</p>{barras}
<h2>3. Situaciones que requieren atención</h2><p>{html.escape(t['incidencias'])}</p>
{ctx}
<p>{html.escape(t['restituidos'])}</p>
<table><thead><tr><th style="width:7%">Hora</th><th style="width:20%">Distrito</th><th style="width:11%">Restituidos RENIEC</th><th style="width:11%">Restituidos JNE</th><th>Qué se reportó</th></tr></thead><tbody>{filas}</tbody></table>
<h2>4. Próximos pasos</h2><p>El corte de cierre de la orientación está previsto hacia las 16:30. Se continuará el seguimiento de los distritos con incidencias graves en coordinación con la PNP, el Ministerio Público y las subprefecturas.</p>
<p class="pie">Fuente: formulario de orientadores, matriz de seguimiento de la SDPEG, información de campo y base de ciudadanos restituidos al 25/09/2026. Datos al {d['corte_datos']}.</p>
</main></body></html>"""


def caja(inf, rotulo: str, texto: str) -> None:
    """Recuadro de texto justificado, con franja roja a la izquierda."""
    par = inf.p("", after=8)
    r = par.add_run(rotulo + " ")
    r.bold = True
    r.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
    par.add_run(texto)
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    izq = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "24"), ("space", "8"), ("color", "B42318")):
        izq.set(qn(f"w:{k}"), v)
    bdr.append(izq)
    shd = OxmlElement("w:shd")
    for k, v in (("val", "clear"), ("color", "auto"), ("fill", "FDF2F1")):
        shd.set(qn(f"w:{k}"), v)
    pPr.insert(0, shd)
    pPr.insert(0, bdr)


def word(d: dict, salida: Path) -> None:
    inf = iw.Informe(d, "INFORME EJECUTIVO DE LA ORIENTACIÓN ELECTORAL", "Documento de uso interno.")
    t, cs = textos(d), d["consultas"]
    k = kpis(d)
    inf.h1("1. Situación general")
    inf.p(t["general"])
    inf.kpis(k[:3])
    inf.kpis(k[3:])
    inf.h1("2. Orientación brindada")
    inf.p(t["orientacion"])
    inf.tabla(["Motivo de la consulta", "Personas", "Porcentaje"], [(m, n, pc(n, cs["total"])) for m, n in cs["por_tipo"].items() if n > 0], [9.0, 3.5, 3.5], centradas=(1, 2))
    inf.h1("3. Situaciones que requieren atención")
    inf.p(t["incidencias"])
    for x in d["contexto"]:
        caja(inf, x["distritos"] + ".", x["texto"])
    inf.p(t["restituidos"])
    inf.tabla(["Hora", "Distrito", "Restituidos RENIEC", "Restituidos JNE", "Qué se reportó"],
              [(x["hora"], f"{x['distrito']} ({dep(x['departamento'])})", fn(x["rest_reniec"]) if x["rest"] else "—", fn(x["rest_jne"]) if x["rest"] else "—", x["texto"]) for x in d["graves"]],
              [1.3, 3.2, 1.9, 1.9, 8.5], centradas=(0, 2, 3))
    inf.h1("4. Próximos pasos")
    inf.p("El corte de cierre de la orientación está previsto hacia las 16:30. Se continuará el seguimiento de los distritos con incidencias graves en coordinación con la PNP, "
          "el Ministerio Público y las subprefecturas.")
    inf.guardar(salida, "Informe ejecutivo", "formulario de orientadores, matriz de seguimiento de la SDPEG, información de campo y base de ciudadanos restituidos al 25/09/2026")


def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    ruta = Path(args[0]) if args else c.ENTRADA
    if "--descargar" in sys.argv:
        url = (c.RAIZ / "data" / "entrada" / "forms_url.txt").read_text(encoding="utf-8").strip()
        urllib.request.urlretrieve(f"{url}&cb={random.randint(0, 10**9)}", ruta)
    d = datos(ruta)
    ic.SEG.mkdir(exist_ok=True)
    if "--word" not in sys.argv:
        (ic.SEG / "boceto_informe_ejecutivo.html").write_text(boceto(d), encoding="utf-8")
        print(f"boceto_informe_ejecutivo.html | datos al {d['corte_datos']} | consultas {d['consultas']['total']} | incidencias {len(d['todas'])} ({len(d['graves'])} graves)")
    if "--boceto" not in sys.argv:
        salida = ic.SEG / f"INFORME_EJECUTIVO_Q20261004_{d['hora'].replace(':', '')}.docx"
        word(d, salida)
        print(f"{salida.name}")


if __name__ == "__main__":
    main()
