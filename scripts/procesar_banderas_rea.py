"""Banderas del dia: denuncias REA informadas (items 75 a 88) para el dashboard de monitoreo.

Fuente (solo lectura): DENUNCIAS_REA/entregables/OBSERVACIONES_INFORMADAS_Q20261002.xlsx, hoja OBSERVACIONES.
Reglas: INSTRUCCIONES_CARGA_OTRO_MAPA.md (misma carpeta). Salida: data/banderas_rea_Q20261002.json.

Lo que no esta en la fila se deja vacio; el texto publicado no se trunca ni se reescribe. Los campos internos
(documento fuente, observacion registrada, supuestos) no salen. UBIGEO INEI siempre como texto de 6 digitos.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ.parent / "DENUNCIAS_REA" / "entregables" / "OBSERVACIONES_INFORMADAS_Q20261002.xlsx"
LOOKUP = RAIZ.parent / "DENUNCIAS_REA" / "data" / "geodata" / "distritos_lookup.parquet"
SALIDA = RAIZ / "data" / "banderas_rea_Q20261002.json"
CORTE = "02/10/2026"
FECHA_ASUMIDA = {80, 82, 83, 86, 87, 88}  # sin sello de RENIEC (instrucciones, regla 6)

CATEGORIAS = {
    "Impugnación o depuración de electores": 1,
    "Verificación o fiscalización de domicilio": 2,
    "Solicitud de padrón o información": 3,
    "Denuncia o alerta preventiva de golondrinaje o transhumancia": 4,
    "Reclamo posterior a verificación o restitución": 5,
}


def leer(ruta: Path = FUENTE) -> list[dict]:
    ws = openpyxl.load_workbook(ruta, data_only=True)["OBSERVACIONES"]
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def construir(filas: list[dict], catalogo: pd.DataFrame) -> list[dict]:
    por_ubigeo = {str(r.UBIGEO_INEI): r for r in catalogo.itertuples()}
    out = []
    for f in filas:
        item = int(f["ITEM"])
        ubigeo = str(f["UBIGEO_INEI"]).strip().zfill(6)
        if not re.fullmatch(r"\d{6}", ubigeo) or ubigeo not in por_ubigeo:
            raise SystemExit(f"Item {item}: ubigeo {ubigeo!r} no existe en el catalogo INEI")
        cat = CATEGORIAS.get(str(f["CATEGORIA_NUEVA"]).strip())
        if cat is None:
            raise SystemExit(f"Item {item}: categoria desconocida {f['CATEGORIA_NUEVA']!r}")
        texto = str(f["OBSERVACION_PUBLICADA_EN_EL_SITIO"] or "").strip()
        if re.search(r"\d{8,}", texto):
            raise SystemExit(f"Item {item}: el texto publicado trae un numero largo (posible DNI). Se detiene la carga.")
        bandera = str(f["BANDERA_POSIBLE_CONFLICTO_O_VIOLENCIA"]).strip().upper() == "SI"
        c = por_ubigeo[ubigeo]
        listados = f["CIUDADANOS_LISTADOS"]
        out.append({
            "id": f"REA-{item}", "item": item, "fecha_recepcion": str(f["FECHA_DE_INGRESO_REGISTRADA"]).strip(),
            "fecha_asumida": item in FECHA_ASUMIDA,
            "departamento": str(c.DEPARTAMEN), "provincia": str(c.PROVINCIA), "distrito": str(c.DISTRITO),
            "ubigeo_inei": ubigeo, "lon": round(float(c.lon), 5), "lat": round(float(c.lat), 5),
            "categoria_rea": cat, "conflictividad": bandera,
            "motivo_conflictividad": str(f["MOTIVO_DE_LA_BANDERA"] or "").strip(),
            "fuente_dashboard": "conflictividad" if bandera else "otro_canal",
            "resumen_publicable": texto, "n_ciudadanos": int(listados) if listados not in (None, "") else None,
            "entidad_emisora": str(f["EMISOR_INSTITUCIONAL"]).strip(), "canal": str(f["CANAL"]).strip(),
            "estado": "VALIDADO", "simulado": False,
        })
    return out


def main() -> None:
    filas = leer()
    cat = pd.read_parquet(LOOKUP, columns=["UBIGEO_INEI", "DEPARTAMEN", "PROVINCIA", "DISTRITO", "lon", "lat"])
    reg = construir(filas, cat)
    con = [r for r in reg if r["conflictividad"]]
    esperado = (len(reg), len(con), len({r["ubigeo_inei"] for r in con}), sum(r["n_ciudadanos"] or 0 for r in reg))
    if esperado != (14, 11, 9, 380):
        raise SystemExit(f"La carga no cuadra con las instrucciones (filas, con bandera, distritos, listados): {esperado} != (14, 11, 9, 380)")
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps({"fuente": "REA, denuncias informadas (items 75 a 88)", "corte": CORTE, "registros": reg},
                                 ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(reg)} filas, {len(con)} con bandera en {len({r['ubigeo_inei'] for r in con})} distritos, "
          f"{esperado[3]} ciudadanos listados -> {SALIDA}")


if __name__ == "__main__":
    sys.exit(main())
